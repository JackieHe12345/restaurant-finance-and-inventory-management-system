from __future__ import annotations

import os
import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

os.environ.setdefault("DATABASE_URL", "sqlite+pysqlite://")
os.environ.setdefault("SECRET_KEY", "dataset-api-test-secret")
os.environ.setdefault("ADMIN_USERNAME", "admin")
os.environ.setdefault("ADMIN_PASSWORD", "Admin123!")
os.environ.setdefault("TEST_LOGIN_ENABLED", "true")
os.environ.setdefault("TEST_DATASET_ID", "test-demo-v1")

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.pool import StaticPool


@compiles(JSONB, "sqlite")
def compile_jsonb_for_sqlite(_type, _compiler, **_kwargs):
    return "JSON"


from app.money import money_text, normalize_payload
from app import database, main
from app.excel_io import WorkbookRecord, build_user_workbook, payload_checksum
from app.models import AuditLog, Dataset, Device, EntryTiming, SyncConflict, SyncRecord, utcnow


def record(record_type: str, key: str, payload: dict) -> WorkbookRecord:
    return WorkbookRecord(record_type, key, payload, payload_checksum(payload))


class DatasetApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = create_engine(
            "sqlite+pysqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )

        @event.listens_for(cls.engine, "connect")
        def enable_sqlite_foreign_keys(connection, _record):
            cursor = connection.cursor()
            cursor.execute("PRAGMA foreign_keys=ON")
            cursor.close()

        database.engine = cls.engine
        database.SessionLocal.configure(bind=cls.engine)
        main.engine = cls.engine
        cls.client_context = TestClient(main.app)
        cls.client = cls.client_context.__enter__()
        login = cls.client.post(
            "/api/v1/auth/login",
            json={"username": "admin", "password": "Admin123!"},
        )
        assert login.status_code == 200, login.text
        cls.headers = {
            "Authorization": f"Bearer {login.json()['access_token']}"
        }

    @classmethod
    def tearDownClass(cls):
        cls.client_context.__exit__(None, None, None)
        cls.engine.dispose()

    def create_dataset(self, dataset_id: str, name: str) -> None:
        response = self.client.post(
            "/api/v1/datasets",
            headers=self.headers,
            json={"id": dataset_id, "name": name, "schema_version": 2},
        )
        self.assertEqual(response.status_code, 200, response.text)

    def put_record(self, dataset_id: str, record_type: str, key: str, payload: dict):
        response = self.client.put(
            f"/api/v1/datasets/{dataset_id}/records/{record_type}/{key}",
            headers=self.headers,
            json={"payload": payload},
        )
        self.assertEqual(response.status_code, 200, response.text)

    def create_role_headers(self, dataset_id: str, username: str, role: str) -> dict[str, str]:
        created = self.client.post(
            "/api/v1/users",
            headers=self.headers,
            json={
                "username": username,
                "password": "Entry123!",
                "role": role,
                "dataset_ids": [dataset_id],
                "default_dataset_id": dataset_id,
            },
        )
        self.assertEqual(created.status_code, 200, created.text)
        login = self.client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": "Entry123!"},
        )
        self.assertEqual(login.status_code, 200, login.text)
        return {"Authorization": f"Bearer {login.json()['access_token']}"}

    def test_admin_changes_are_written_to_audit_log(self):
        dataset_id = "admin-audit-ledger"
        self.create_dataset(dataset_id, "管理员记录测试")
        self.put_record(dataset_id, "global", "admin-change", {"value": 1})
        with database.SessionLocal() as db:
            audit = db.scalar(
                select(AuditLog)
                .where(AuditLog.dataset_id == dataset_id)
                .order_by(AuditLog.created_at.desc())
            )
            self.assertIsNotNone(audit)
            self.assertEqual(audit.actor_username, "admin")
            self.assertEqual(audit.record_key, "admin-change")

    def test_admin_can_rename_and_safely_delete_unbound_dataset(self):
        dataset_id = "rename-delete-ledger"
        self.create_dataset(dataset_id, "原名称")
        renamed = self.client.patch(
            f"/api/v1/datasets/{dataset_id}",
            headers=self.headers,
            json={"name": " 新名称 "},
        )
        self.assertEqual(renamed.status_code, 200, renamed.text)
        self.assertEqual(renamed.json()["name"], "新名称")

        self.put_record(dataset_id, "global", "orphan-check", {"value": 1})
        device = self.client.post(
            "/api/v1/devices/register",
            headers=self.headers,
            json={"dataset_id": dataset_id, "name": "test-device"},
        )
        self.assertEqual(device.status_code, 200, device.text)
        with database.SessionLocal() as db:
            db.add(
                SyncConflict(
                    dataset_id=dataset_id,
                    record_type="global",
                    record_key="orphan-check",
                )
            )
            db.commit()

        mismatch = self.client.delete(
            f"/api/v1/datasets/{dataset_id}?confirm_name=原名称",
            headers=self.headers,
        )
        self.assertEqual(mismatch.status_code, 409, mismatch.text)

        deleted = self.client.delete(
            f"/api/v1/datasets/{dataset_id}?confirm_name=新名称",
            headers=self.headers,
        )
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertEqual(
            deleted.json(),
            {"deleted": True, "id": dataset_id, "name": "新名称"},
        )
        with database.SessionLocal() as db:
            self.assertIsNone(db.get(Dataset, dataset_id))
            self.assertEqual(
                db.scalar(
                    select(func.count())
                    .select_from(SyncRecord)
                    .where(SyncRecord.dataset_id == dataset_id)
                ),
                0,
            )
            self.assertEqual(
                db.scalar(
                    select(func.count())
                    .select_from(Device)
                    .where(Device.dataset_id == dataset_id)
                ),
                0,
            )
            self.assertEqual(
                db.scalar(
                    select(func.count())
                    .select_from(AuditLog)
                    .where(AuditLog.dataset_id == dataset_id)
                ),
                0,
            )
            self.assertEqual(
                db.scalar(
                    select(func.count())
                    .select_from(SyncConflict)
                    .where(SyncConflict.dataset_id == dataset_id)
                ),
                0,
            )

    def test_delete_rejects_test_and_bound_datasets(self):
        test_delete = self.client.delete(
            "/api/v1/datasets/test-demo-v1?confirm_name=Test%20独立演示账本",
            headers=self.headers,
        )
        self.assertEqual(test_delete.status_code, 422, test_delete.text)

        dataset_id = "bound-ledger"
        self.create_dataset(dataset_id, "已绑定账本")
        created_user = self.client.post(
            "/api/v1/users",
            headers=self.headers,
            json={
                "username": "bound-user",
                "password": "Editor123!",
                "role": "editor",
                "dataset_ids": [dataset_id],
                "default_dataset_id": dataset_id,
            },
        )
        self.assertEqual(created_user.status_code, 200, created_user.text)
        response = self.client.delete(
            f"/api/v1/datasets/{dataset_id}?confirm_name=已绑定账本",
            headers=self.headers,
        )
        self.assertEqual(response.status_code, 409, response.text)
        self.assertIn("解除绑定", response.json()["detail"])

    def test_fixed_expense_app_checkbox_wins_over_stale_web_alias(self):
        dataset_id = "canonical-fixed-cash-flag"
        self.create_dataset(dataset_id, "固定成本同步")
        saved = self.client.put(
            f"/api/v1/datasets/{dataset_id}/records/day/2026-09-02",
            headers=self.headers,
            json={
                "payload": {
                    "fixed_expenses": [
                        {
                            "日期": "2026-09-02",
                            "类别": "房租",
                            "金额": 500,
                            # An App save must beat a stale cached web value.
                            "计入可用余额": False,
                            "include_cash_balance": True,
                        }
                    ]
                }
            },
        )
        self.assertEqual(saved.status_code, 200, saved.text)
        expense = saved.json()["payload"]["fixed_expenses"][0]
        self.assertIs(expense["计入可用余额"], False)
        self.assertNotIn("include_cash_balance", expense)

        cached = self.client.get(
            f"/api/v1/datasets/{dataset_id}/records/day/2026-09-02",
            headers=self.headers,
        )
        self.assertEqual(cached.status_code, 200, cached.text)
        self.assertIs(cached.json()["payload"]["fixed_expenses"][0]["计入可用余额"], False)
        with database.SessionLocal() as db:
            flow = main._cash_flow_summary_db(
                db,
                dataset_id,
                {},
                now=datetime(2026, 9, 3, tzinfo=timezone.utc),
            )
        self.assertEqual(flow["fixed_expenses"], money_text(0))

    def test_internal_audit_cache_deletes_immediately(self):
        dataset_id = "internal-audit-cache"
        self.create_dataset(dataset_id, "内部修改记录")
        self.put_record(dataset_id, "global", "audit-entry", {"value": 1})
        audits = self.client.get(
            f"/api/v1/audit?dataset_id={dataset_id}", headers=self.headers
        )
        self.assertEqual(audits.status_code, 200, audits.text)
        audit_id = audits.json()[0]["id"]

        exported = self.client.get(
            f"/api/v1/datasets/{dataset_id}/workbooks/internal", headers=self.headers
        )
        self.assertEqual(exported.status_code, 200, exported.text)
        workbook = load_workbook(BytesIO(exported.content), data_only=True)
        self.assertEqual(workbook["__audit_records"].sheet_state, "veryHidden")
        exported_ids = {
            row[0]
            for row in workbook["__audit_records"].iter_rows(
                min_row=2, min_col=2, max_col=2, values_only=True
            )
            if row[0]
        }
        workbook.close()
        self.assertIn(audit_id, exported_ids)

        deleted = self.client.delete(f"/api/v1/audit/{audit_id}", headers=self.headers)
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertEqual(deleted.json()["deleted_count"], 1)
        refreshed = self.client.get(
            f"/api/v1/datasets/{dataset_id}/workbooks/internal", headers=self.headers
        )
        self.assertEqual(refreshed.status_code, 200, refreshed.text)
        workbook = load_workbook(BytesIO(refreshed.content), data_only=True)
        refreshed_ids = {
            row[0]
            for row in workbook["__audit_records"].iter_rows(
                min_row=2, min_col=2, max_col=2, values_only=True
            )
            if row[0]
        }
        workbook.close()
        self.assertNotIn(audit_id, refreshed_ids)

    def test_user_workbook_hides_and_preserves_fixed_cash_checkbox(self):
        dataset_id = "user-workbook-hidden-cash-flag"
        self.create_dataset(dataset_id, "用户 Excel 内部开关")
        self.put_record(
            dataset_id,
            "day",
            "2026-09-02",
            {
                "fixed_expenses": [
                    {
                        "日期": "2026-09-02",
                        "类别": "房租",
                        "明细": "九月房租",
                        "金额": 500,
                        "计入可用余额": False,
                    }
                ]
            },
        )
        exported = self.client.get(
            f"/api/v1/datasets/{dataset_id}/workbooks/user?year=2026",
            headers=self.headers,
        )
        self.assertEqual(exported.status_code, 200, exported.text)
        workbook = load_workbook(BytesIO(exported.content), data_only=True)
        headers = [cell.value for cell in workbook["fixed_expenses"][1]]
        workbook.close()
        self.assertNotIn("计入可用余额", headers)
        self.assertNotIn("include_cash_balance", headers)

        uploaded = self.client.post(
            f"/api/v1/datasets/{dataset_id}/workbooks/user?year=2026",
            headers=self.headers,
            files={
                "file": (
                    "user.xlsx",
                    exported.content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        record_response = self.client.get(
            f"/api/v1/datasets/{dataset_id}/records/day/2026-09-02",
            headers=self.headers,
        )
        self.assertEqual(record_response.status_code, 200, record_response.text)
        expense = record_response.json()["payload"]["fixed_expenses"][0]
        self.assertIs(expense["计入可用余额"], False)

    def test_admin_can_overwrite_test_internal_data_without_renaming_it(self):
        source_dataset_id = "source-for-test-overwrite"
        self.create_dataset(source_dataset_id, "覆盖来源账本")
        self.put_record(source_dataset_id, "global", "source-value", {"value": 7})
        source_workbook = self.client.get(
            f"/api/v1/datasets/{source_dataset_id}/workbooks/internal",
            headers=self.headers,
        )
        self.assertEqual(source_workbook.status_code, 200, source_workbook.text)

        with TemporaryDirectory() as temporary_dir:
            with patch.object(main, "Path", return_value=Path(temporary_dir)):
                overwritten = self.client.post(
                    "/api/v1/datasets/test-demo-v1/workbooks/internal-overwrite",
                    headers=self.headers,
                    data={
                        "confirm_target_name": "Test 独立演示账本",
                        "confirm_again": "Test 独立演示账本",
                    },
                    files={
                        "file": (
                            "source-internal.xlsx",
                            source_workbook.content,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        )
                    },
                )
        self.assertEqual(overwritten.status_code, 200, overwritten.text)
        self.assertEqual(overwritten.json()["target_dataset_name"], "Test 独立演示账本")
        with database.SessionLocal() as db:
            target = db.get(Dataset, "test-demo-v1")
            self.assertIsNotNone(target)
            self.assertEqual(target.name, "Test 独立演示账本")
            copied = db.scalar(
                select(SyncRecord).where(
                    SyncRecord.dataset_id == "test-demo-v1",
                    SyncRecord.record_type == "global",
                    SyncRecord.record_key == "source-value",
                    SyncRecord.is_deleted.is_(False),
                )
            )
            self.assertIsNotNone(copied)
            self.assertEqual(copied.payload, {"value": 7})

    def test_user_workbook_replaces_only_selected_year_and_two_globals(self):
        dataset_id = "replace-year-ledger"
        self.create_dataset(dataset_id, "整年替换账本")
        self.put_record(dataset_id, "day", "2025-12-31", {"keep": "other-year"})
        self.put_record(
            dataset_id,
            "day",
            "2026-01-01",
            {"daily": [{"日期": "2026-01-01", "收益": 1}], "legacy": "remove"},
        )
        self.put_record(dataset_id, "day", "2026-02-02", {"daily": []})
        self.put_record(dataset_id, "month", "2026-01", {"monthly_fixed": []})
        self.put_record(dataset_id, "month", "2026-02", {"monthly_fixed": []})
        self.put_record(
            dataset_id,
            "global",
            "ingredients",
            {"ingredients": [{"词条名称": "旧食材", "分类": "旧"}]},
        )
        self.put_record(
            dataset_id,
            "global",
            "fixed_categories",
            {"fixed_categories": [{"类别名称": "旧类别"}]},
        )
        self.put_record(dataset_id, "global", "custom-global", {"keep": True})

        content = build_user_workbook(
            dataset_id=dataset_id,
            schema_version=2,
            year=2026,
            records=[
                record(
                    "day",
                    "2026-01-01",
                    {
                        "daily": [{"日期": "2026-01-01", "收益": 99}],
                        "fixed_expenses": [],
                        "dishes": [],
                    },
                ),
                record(
                    "month",
                    "2026-01",
                    {"monthly_fixed": [{"月份": "2026-01", "人工总额": 8}]},
                ),
                record(
                    "global",
                    "ingredients",
                    {"ingredients": [{"词条名称": "新食材", "分类": "新"}]},
                ),
                record(
                    "global",
                    "fixed_categories",
                    {"fixed_categories": []},
                ),
            ],
        )
        upload = self.client.post(
            f"/api/v1/datasets/{dataset_id}/workbooks/user?year=2026",
            headers=self.headers,
            files={
                "file": (
                    "replacement.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(upload.status_code, 200, upload.text)
        result = upload.json()
        self.assertEqual(result["mode"], "replace_year")
        self.assertEqual(result["replace_year"], 2026)
        self.assertEqual(result["deleted_records"], 2)

        records_response = self.client.get(
            f"/api/v1/datasets/{dataset_id}/records?include_deleted=true&limit=100",
            headers=self.headers,
        )
        self.assertEqual(records_response.status_code, 200, records_response.text)
        by_identity = {
            (item["record_type"], item["record_key"]): item
            for item in records_response.json()
        }
        self.assertFalse(by_identity[("day", "2025-12-31")]["is_deleted"])
        replaced_day = by_identity[("day", "2026-01-01")]
        self.assertFalse(replaced_day["is_deleted"])
        self.assertNotIn("legacy", replaced_day["payload"])
        self.assertEqual(replaced_day["payload"]["daily"][0]["收益"], money_text(99))
        self.assertTrue(by_identity[("day", "2026-02-02")]["is_deleted"])
        self.assertTrue(by_identity[("month", "2026-02")]["is_deleted"])
        self.assertFalse(by_identity[("global", "custom-global")]["is_deleted"])
        self.assertEqual(
            by_identity[("global", "ingredients")]["payload"]["ingredients"][0][
                "词条名称"
            ],
            "新食材",
        )
        self.assertEqual(
            by_identity[("global", "fixed_categories")]["payload"],
            {"fixed_categories": []},
        )

    def test_non_admin_cannot_upload_user_workbook(self):
        dataset_id = "editor-download-only"
        self.create_dataset(dataset_id, "普通账号只允许下载")
        created_user = self.client.post(
            "/api/v1/users",
            headers=self.headers,
            json={
                "username": "download-only-user",
                "password": "Editor123!",
                "role": "editor",
                "dataset_ids": [dataset_id],
                "default_dataset_id": dataset_id,
            },
        )
        self.assertEqual(created_user.status_code, 200, created_user.text)
        login = self.client.post(
            "/api/v1/auth/login",
            json={"username": "download-only-user", "password": "Editor123!"},
        )
        self.assertEqual(login.status_code, 200, login.text)
        editor_headers = {
            "Authorization": f"Bearer {login.json()['access_token']}"
        }
        content = build_user_workbook(
            dataset_id=dataset_id,
            schema_version=2,
            year=2026,
            records=[],
        )

        download = self.client.get(
            f"/api/v1/datasets/{dataset_id}/workbooks/user?year=2026",
            headers=editor_headers,
        )
        self.assertEqual(download.status_code, 200, download.text)
        upload = self.client.post(
            f"/api/v1/datasets/{dataset_id}/workbooks/user?year=2026",
            headers=editor_headers,
            files={
                "file": (
                    "editor-upload.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(upload.status_code, 403, upload.text)
        merge_upload = self.client.post(
            f"/api/v1/datasets/{dataset_id}/workbooks/user-merge?year=2026",
            headers=editor_headers,
            files={
                "file": (
                    "editor-merge.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(merge_upload.status_code, 403, merge_upload.text)

    def test_admin_merge_user_workbook_adds_quantities_and_keeps_unit_price(self):
        dataset_id = "merge-workbook-ledger"
        self.create_dataset(dataset_id, "合并账本")
        existing_payload = {
            "daily": [{
                "日期": "2026-03-08",
                "收益": 100,
                "新客人数量": 2,
                "总客人数": 5,
                "食材成本": 40,
                "食材明细": "牛肉 2斤 40",
                "佐料成本": 0,
                "佐料明细": "",
                "水电": 0,
                "其他": 0,
                "其他明细": "",
                "当日备注": "第一批",
            }],
            "fixed_expenses": [{"日期": "2026-03-08", "类别": "运费", "明细": "配送", "金额": 5}],
            "dishes": [],
        }
        self.put_record(dataset_id, "day", "2026-03-08", existing_payload)
        self.put_record(
            dataset_id,
            "global",
            "ingredients",
            {"ingredients": [{"词条名称": "牛肉", "分类": "食材"}]},
        )
        content = build_user_workbook(
            dataset_id=dataset_id,
            schema_version=2,
            year=2026,
            records=[
                record(
                    "day",
                    "2026-03-08",
                    {
                        "daily": [{
                            "日期": "2026-03-08",
                            "收益": 100,
                            "新客人数量": 2,
                            "总客人数": 5,
                            "食材成本": 40,
                            "食材明细": "牛肉 2斤 40",
                            "佐料成本": 0,
                            "佐料明细": "",
                            "水电": 0,
                            "其他": 0,
                            "其他明细": "",
                            "当日备注": "第二批",
                        }],
                        "fixed_expenses": [{"日期": "2026-03-08", "类别": "运费", "明细": "配送", "金额": 5}],
                        "dishes": [],
                    },
                ),
                record(
                    "global",
                    "ingredients",
                    {"ingredients": [{"词条名称": "牛肉", "分类": "食材"}, {"词条名称": "土豆", "分类": "食材"}]},
                ),
            ],
        )
        upload = self.client.post(
            f"/api/v1/datasets/{dataset_id}/workbooks/user-merge?year=2026",
            headers=self.headers,
            files={
                "file": (
                    "merge.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(upload.status_code, 200, upload.text)
        self.assertEqual(upload.json()["mode"], "merge")
        records = self.client.get(
            f"/api/v1/datasets/{dataset_id}/records?limit=100",
            headers=self.headers,
        ).json()
        day = next(item for item in records if item["record_type"] == "day")
        daily = day["payload"]["daily"][0]
        self.assertEqual(daily["收益"], money_text(200))
        self.assertEqual(daily["食材成本"], money_text(80))
        self.assertEqual(daily["食材明细"], "牛肉 4斤 80.00")
        self.assertEqual(len(day["payload"]["fixed_expenses"]), 2)
        ingredients = next(item for item in records if item["record_key"] == "ingredients")
        self.assertEqual(len(ingredients["payload"]["ingredients"]), 2)

    def test_entry_nine_income_slots_are_independent_and_use_11am_cutoff(self):
        dataset_id = "sep01-entry-income-nine"
        self.create_dataset(dataset_id, "Sep01 九个收入时间窗")
        entry_headers = self.create_role_headers(
            dataset_id, "sep01-entry-income-user", "entry"
        )
        day = "2026-09-02"
        url = f"/api/v1/datasets/{dataset_id}/records/day/{day}"

        # 10:00: A cash starts only its own clock.
        row = {"日期": day, "A店现金": 100}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 2, 0, tzinfo=timezone.utc)):
            saved = self.client.put(url, headers=entry_headers, json={"payload": {"daily": [row]}})
        self.assertEqual(saved.status_code, 200, saved.text)

        # 10:31: the still-empty A WeChat slot remains independently writable;
        # saving zero is a real first save and therefore starts its clock.
        row = {**row, "A店微信覆盖": 0}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 2, 31, tzinfo=timezone.utc)):
            saved_zero = self.client.put(url, headers=entry_headers, json={"payload": {"daily": [row]}})
        self.assertEqual(saved_zero.status_code, 200, saved_zero.text)

        changed_cash = {**row, "A店现金": 101}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 2, 31, tzinfo=timezone.utc)):
            cash_locked = self.client.put(url, headers=entry_headers, json={"payload": {"daily": [changed_cash]}})
        self.assertEqual(cash_locked.status_code, 403, cash_locked.text)

        changed_zero = {**row, "A店微信覆盖": 5}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 3, 1, tzinfo=timezone.utc)):
            zero_locked = self.client.put(url, headers=entry_headers, json={"payload": {"daily": [changed_zero]}})
        self.assertEqual(zero_locked.status_code, 403, zero_locked.text)

        # The student-delivery slot can be first-saved at 10:59 next day. Its
        # own 30-minute edit window survives the 11:00 first-entry cutoff.
        row = {**row, "学生外卖": 50}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 3, 2, 59, tzinfo=timezone.utc)):
            late_first = self.client.put(url, headers=entry_headers, json={"payload": {"daily": [row]}})
        self.assertEqual(late_first.status_code, 200, late_first.text)

        at_cutoff = {**row, "B店外卖": 1}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 3, 3, 0, tzinfo=timezone.utc)):
            cutoff_rejected = self.client.put(url, headers=entry_headers, json={"payload": {"daily": [at_cutoff]}})
        self.assertEqual(cutoff_rejected.status_code, 403, cutoff_rejected.text)
        self.assertIn("11:00", cutoff_rejected.text)

        row = {**row, "学生外卖": 60}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 3, 3, 20, tzinfo=timezone.utc)):
            after_cutoff_edit = self.client.put(url, headers=entry_headers, json={"payload": {"daily": [row]}})
            windows_response = self.client.get(
                f"/api/v1/datasets/{dataset_id}/entry-windows", headers=entry_headers
            )
        self.assertEqual(after_cutoff_edit.status_code, 200, after_cutoff_edit.text)
        self.assertEqual(windows_response.status_code, 200, windows_response.text)
        window = next(item for item in windows_response.json() if item["record_key"] == day)
        self.assertEqual(set(window["income_fields"]), set(main.ENTRY_INCOME_SLOTS))
        self.assertTrue(window["income_fields"]["a_cash"]["saved_at"].startswith("2026-09-02T02:00"))
        self.assertTrue(window["income_fields"]["a_wechat"]["saved_at"].startswith("2026-09-02T02:31"))
        self.assertTrue(window["income_fields"]["student_takeout"]["saved_at"].startswith("2026-09-03T02:59"))
        self.assertEqual(window["income_fields"]["b_takeout"]["state"], "expired")

    def test_entry_detail_items_have_independent_30m_and_category_12h_windows(self):
        dataset_id = "sep01-entry-detail-items"
        self.create_dataset(dataset_id, "Sep01 明细项目时间窗")
        entry_headers = self.create_role_headers(
            dataset_id, "sep01-entry-detail-user", "entry"
        )
        day = "2026-09-02"
        url = f"/api/v1/datasets/{dataset_id}/records/day/{day}"

        def save(local_hour: int, local_minute: int, current_row: dict):
            utc_hour = local_hour - 8
            stamp_day = 2
            if utc_hour < 0:
                utc_hour += 24
                stamp_day -= 1
            with patch.object(
                main,
                "utcnow",
                return_value=datetime(2026, 9, stamp_day, utc_hour, local_minute, tzinfo=timezone.utc),
            ):
                return self.client.put(
                    url, headers=entry_headers, json={"payload": {"daily": [current_row]}}
                )

        # 10:00: six food items start only food's 12-hour addition window,
        # plus six independent 30-minute item clocks.
        food = "土豆 10, 猪肉 20, 青菜 30, 鸡蛋 40, 大米 50, 油 60"
        row = {"日期": day, "食材明细": food, "食材成本": 210}
        self.assertEqual(save(10, 0, row).status_code, 200)

        # One original item is editable before 10:30, but another original item
        # is locked at 10:31. A successful correction does not renew its clock.
        food = "土豆 11, 猪肉 20, 青菜 30, 鸡蛋 40, 大米 50, 油 60"
        row = {**row, "食材明细": food, "食材成本": 211}
        self.assertEqual(save(10, 20, row).status_code, 200)
        expired_original = {**row, "食材明细": food.replace("猪肉 20", "猪肉 21")}
        self.assertEqual(save(10, 31, expired_original).status_code, 403)

        # Items appended later receive their own first-save clocks.
        food += ", 牛肉 100"
        row = {**row, "食材明细": food, "食材成本": 311}
        self.assertEqual(save(19, 0, row).status_code, 200)
        food = food.replace("牛肉 100", "牛肉 110")
        row = {**row, "食材明细": food, "食材成本": 321}
        self.assertEqual(save(19, 20, row).status_code, 200)
        late_beef = {**row, "食材明细": food.replace("牛肉 110", "牛肉 120")}
        self.assertEqual(save(19, 30, late_beef).status_code, 403)

        food += ", 白切鸡 80"
        row = {**row, "食材明细": food, "食材成本": 401}
        self.assertEqual(save(20, 0, row).status_code, 200)

        row = {**row, "佐料明细": "酱油 10", "佐料成本": 10}
        self.assertEqual(save(21, 0, row).status_code, 200)
        row = {**row, "佐料明细": "酱油 10, 食用盐 5", "佐料成本": 15}
        self.assertEqual(save(21, 30, row).status_code, 200)
        late_soy = {**row, "佐料明细": "酱油 11, 食用盐 5", "佐料成本": 16}
        self.assertEqual(save(21, 31, late_soy).status_code, 403)

        row = {**row, "其他明细": "Boss直聘 100", "其他": 100}
        self.assertEqual(save(21, 40, row).status_code, 200)
        row = {**row, "其他明细": "Boss直聘 120", "其他": 120}
        self.assertEqual(save(22, 0, row).status_code, 200)

        # Exactly 12 hours after food's first save, only food is closed. Spice
        # and other keep their own later 12-hour deadlines.
        add_food = {**row, "食材明细": food + ", 蘑菇 10"}
        cutoff = save(22, 0, add_food)
        self.assertEqual(cutoff.status_code, 403, cutoff.text)
        self.assertIn("12", cutoff.text)
        row = {**row, "佐料明细": "酱油 10, 食用盐 5, 胡椒 3", "佐料成本": 18}
        self.assertEqual(save(22, 1, row).status_code, 200)
        late_other = {**row, "其他明细": "Boss直聘 130", "其他": 130}
        self.assertEqual(save(22, 10, late_other).status_code, 403)

        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 14, 5, tzinfo=timezone.utc)):
            windows = self.client.get(
                f"/api/v1/datasets/{dataset_id}/entry-windows", headers=entry_headers
            ).json()
        window = next(item for item in windows if item["record_key"] == day)
        self.assertEqual(len(set(window["detail_deadlines"].values())), 3)
        self.assertTrue(window["detail_deadlines"]["food"].startswith("2026-09-02T14:00"))
        self.assertTrue(window["detail_deadlines"]["spice"].startswith("2026-09-03T01:00"))
        self.assertTrue(window["detail_deadlines"]["other"].startswith("2026-09-03T01:40"))
        by_label = {item["label"]: item for item in window["detail_items"]}
        self.assertTrue(by_label["牛肉"]["saved_at"].startswith("2026-09-02T11:00"))
        self.assertTrue(by_label["白切鸡"]["saved_at"].startswith("2026-09-02T12:00"))
        self.assertTrue(by_label["酱油"]["saved_at"].startswith("2026-09-02T13:00"))
        self.assertTrue(by_label["食用盐"]["saved_at"].startswith("2026-09-02T13:30"))
        self.assertTrue(by_label["胡椒"]["saved_at"].startswith("2026-09-02T14:01"))
        self.assertTrue(by_label["Boss直聘"]["saved_at"].startswith("2026-09-02T13:40"))

    def test_entry_detail_rename_keeps_original_edit_clock(self):
        dataset_id = "sep01-entry-detail-rename"
        self.create_dataset(dataset_id, "Sep01 明细改名不重置计时")
        entry_headers = self.create_role_headers(
            dataset_id, "sep01-entry-detail-rename-user", "entry"
        )
        day = "2026-09-02"
        url = f"/api/v1/datasets/{dataset_id}/records/day/{day}"

        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 2, 0, tzinfo=timezone.utc)):
            first = self.client.put(
                url,
                headers=entry_headers,
                json={"payload": {"daily": [{"日期": day, "食材明细": "牛内 80", "食材成本": 80}]}},
            )
        self.assertEqual(first.status_code, 200, first.text)

        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 2, 20, tzinfo=timezone.utc)):
            renamed = self.client.put(
                url,
                headers=entry_headers,
                json={"payload": {"daily": [{"日期": day, "食材明细": "牛肉 80", "食材成本": 80}]}},
            )
        self.assertEqual(renamed.status_code, 200, renamed.text)

        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 2, 31, tzinfo=timezone.utc)):
            late_edit = self.client.put(
                url,
                headers=entry_headers,
                json={"payload": {"daily": [{"日期": day, "食材明细": "牛肉 90", "食材成本": 90}]}},
            )
        self.assertEqual(late_edit.status_code, 403, late_edit.text)

    def test_entry_windows_include_existing_days_without_timing_rows(self):
        dataset_id = "sep01-entry-window-union"
        self.create_dataset(dataset_id, "Sep01 旧数据时间窗")
        day = "2026-09-02"
        self.put_record(dataset_id, "day", day, {"daily": [{"日期": day}]})
        entry_headers = self.create_role_headers(
            dataset_id, "sep01-entry-union-user", "entry"
        )
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 2, 0, tzinfo=timezone.utc)):
            response = self.client.get(
                f"/api/v1/datasets/{dataset_id}/entry-windows", headers=entry_headers
            )
        self.assertEqual(response.status_code, 200, response.text)
        window = next(item for item in response.json() if item["record_key"] == day)
        self.assertEqual(set(window["income_fields"]), set(main.ENTRY_INCOME_SLOTS))
        self.assertTrue(all(item["saved_at"] is None for item in window["income_fields"].values()))
        self.assertTrue(all(item["entry_deadline"].startswith("2026-09-03T03:00") for item in window["income_fields"].values()))

    def test_entry_cannot_bypass_individual_clocks_with_record_delete_or_wechat_sync(self):
        dataset_id = "sep01-entry-no-bypass"
        self.create_dataset(dataset_id, "Sep01 录入权限边界")
        day = "2026-09-02"
        self.put_record(
            dataset_id,
            "day",
            day,
            {"daily": [{"日期": day, "A店现金": 100, "食材明细": "牛肉 80"}]},
        )
        entry_headers = self.create_role_headers(
            dataset_id, "sep01-entry-no-bypass-user", "entry"
        )

        deleted = self.client.post(
            f"/api/v1/datasets/{dataset_id}/records/day/{day}/delete",
            headers=entry_headers,
            json={},
        )
        self.assertEqual(deleted.status_code, 403, deleted.text)

        synced = self.client.post(
            f"/api/v1/datasets/{dataset_id}/wechat/sync?bill_date={day}",
            headers=entry_headers,
        )
        self.assertEqual(synced.status_code, 404, synced.text)

    def test_admin_can_schedule_exact_day_temporary_permission(self):
        dataset_id = "sep04-temporary-permission"
        self.create_dataset(dataset_id, "Sep04 临时权限")
        viewer_headers = self.create_role_headers(
            dataset_id, "sep04-viewer-with-grant", "viewer"
        )
        grant = self.client.post(
            "/api/v1/temporary-permissions",
            headers=self.headers,
            json={
                "user_id": self.client.get("/api/v1/users", headers=self.headers).json()[-1]["id"],
                "dataset_id": dataset_id,
                "permission_role": "editor",
                "business_start_date": "2026-09-04",
                "business_end_date": "2026-09-04",
                "starts_at": "2026-09-04T01:00:00Z",
                "ends_at": "2026-09-04T03:00:00Z",
            },
        )
        self.assertEqual(grant.status_code, 200, grant.text)
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 4, 2, 0, tzinfo=timezone.utc)):
            me = self.client.get("/api/v1/auth/me", headers=viewer_headers)
            allowed = self.client.put(
                f"/api/v1/datasets/{dataset_id}/records/day/2026-09-04",
                headers=viewer_headers,
                json={"payload": {"daily": [{"日期": "2026-09-04", "A店现金": 100}]}},
            )
            wrong_day = self.client.put(
                f"/api/v1/datasets/{dataset_id}/records/day/2026-09-05",
                headers=viewer_headers,
                json={"payload": {"daily": [{"日期": "2026-09-05", "A店现金": 100}]}},
            )
        self.assertEqual(me.status_code, 200, me.text)
        self.assertEqual(len(me.json()["temporary_permissions"]), 1)
        self.assertEqual(allowed.status_code, 200, allowed.text)
        self.assertEqual(wrong_day.status_code, 403, wrong_day.text)
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 4, 3, 0, tzinfo=timezone.utc)):
            expired = self.client.put(
                f"/api/v1/datasets/{dataset_id}/records/day/2026-09-04",
                headers=viewer_headers,
                json={"payload": {"daily": [{"日期": "2026-09-04", "A店现金": 101}]}},
            )
        self.assertEqual(expired.status_code, 403, expired.text)

    def test_cash_position_includes_loan_and_takeout_revenue(self):
        dataset_id = "sep01-cash-position"
        self.create_dataset(dataset_id, "Sep01 可用余额")
        self.put_record(
            dataset_id,
            "day",
            "2026-09-02",
            {
                "daily": [{
                    "日期": "2026-09-02",
                    "A店现金": 300,
                    "B店现金": 200,
                    "学生外卖": 100,
                    "A店外卖": 250,
                    "B店外卖": 150,
                    "食材成本": 100,
                }],
                "fixed_expenses": [{
                    "日期": "2026-09-02",
                    "类别": "房租",
                    "金额": 500,
                    "计入可用余额": False,
                }],
            },
        )
        saved = self.client.put(
            f"/api/v1/datasets/{dataset_id}/cash-position",
            headers=self.headers,
            json={
                "opening_balance": 200,
                "opening_date": "2026-09-01",
                "shareholder_loan": 400,
            },
        )
        self.assertEqual(saved.status_code, 200, saved.text)
        body = saved.json()
        self.assertEqual(body["income"], money_text(1000))
        self.assertEqual(body["total_expenses"], money_text(100))
        self.assertEqual(body["current_balance"], money_text(1500))

        itemized = self.client.put(
            f"/api/v1/datasets/{dataset_id}/cash-position",
            headers=self.headers,
            json={
                "opening_balance": 200,
                "opening_date": "2026-09-01",
                "transactions": [{
                    "id": "loan-on-baseline", "date": "2026-09-01",
                    "type": "向股东借款", "amount": 400, "description": "股东借款",
                }],
            },
        )
        self.assertEqual(itemized.status_code, 200, itemized.text)
        self.assertEqual(itemized.json()["current_balance"], money_text(1500))
        self.assertEqual(itemized.json()["shareholder_loan_total"], money_text(400))

    def test_cash_force_rows_set_balance_and_report_month_year_totals(self):
        dataset_id = "sep10-force-cash-ledger"
        self.create_dataset(dataset_id, "Sep10 现金余额表")
        self.put_record(
            dataset_id,
            "day",
            "2026-09-08",
            {"daily": [{
                "日期": "2026-09-08", "收益": 50, "食材成本": 20,
            }]},
        )
        with patch.object(
            main, "utcnow", return_value=datetime(2026, 9, 10, 4, 0, tzinfo=timezone.utc)
        ):
            saved = self.client.put(
                f"/api/v1/datasets/{dataset_id}/cash-position",
                headers=self.headers,
                json={
                    "opening_balance": 100,
                    "opening_date": "2026-08-31",
                    "transactions": [
                        {
                            "id": "loan", "date": "2026-09-01",
                            "type": "向股东借款", "amount": 200,
                        },
                        {
                            "id": "expense", "date": "2026-09-08",
                            "type": "其它现金支出", "amount": 10,
                        },
                        {
                            "id": "locked-zero", "date": "2026-09-08",
                            "type": "余额强制修改", "amount": 0,
                            "mode": "force", "system_locked": True,
                            "created_at": "2026-09-08T10:00:00+00:00",
                        },
                        {
                            "id": "manual-force", "date": "2026-09-08",
                            "type": "余额强制修改", "amount": 75,
                            "mode": "force", "system_locked": False,
                            "created_at": "2026-09-08T11:00:00+00:00",
                        },
                    ],
                },
            )
        self.assertEqual(saved.status_code, 200, saved.text)
        body = saved.json()
        self.assertEqual(body["current_balance"], money_text(75))
        self.assertEqual(body["previous_month_balance"], money_text(100))
        self.assertEqual(body["shareholder_loan_this_month"], money_text(200))
        self.assertEqual(body["shareholder_loan_this_year"], money_text(200))
        self.assertEqual(body["shareholder_loan_total"], money_text(200))
        self.assertTrue(all("running_balance" in row for row in body["entries"]))
        force_rows = {row["id"]: row for row in body["entries"] if row.get("mode") == "force"}
        self.assertEqual(force_rows["locked-zero"]["income"], None)
        self.assertEqual(force_rows["locked-zero"]["expense"], None)
        self.assertEqual(force_rows["locked-zero"]["running_balance"], money_text(0))
        self.assertFalse(force_rows["locked-zero"]["deletable"])
        self.assertEqual(force_rows["manual-force"]["running_balance"], money_text(75))
        self.assertTrue(force_rows["manual-force"]["deletable"])
        stored = body["record"]["payload"]["transactions"]
        locked = next(item for item in stored if item["id"] == "locked-zero")
        self.assertEqual(
            locked,
            {
                "id": "locked-zero", "date": "2026-09-08",
                "type": "余额强制修改", "amount": "0.00",
                "description": "余额强制修改", "mode": "force",
                "system_locked": True,
                "created_at": "2026-09-08T10:00:00+00:00",
            },
        )

    def test_locked_force_cannot_be_removed_or_changed_by_any_record_route(self):
        dataset_id = "sep10-locked-force-protection"
        self.create_dataset(dataset_id, "Sep10 锁定修正保护")
        original_transactions = [{
            "id": "locked-force", "date": "2026-09-08",
            "type": "余额强制修改", "amount": 0,
            "mode": "force", "system_locked": True,
            "created_at": "2026-09-08T10:00:00Z",
        }]
        created = self.client.put(
            f"/api/v1/datasets/{dataset_id}/cash-position",
            headers=self.headers,
            json={
                "opening_balance": 100,
                "opening_date": "2026-09-01",
                "transactions": original_transactions,
            },
        )
        self.assertEqual(created.status_code, 200, created.text)
        canonical_payload = created.json()["record"]["payload"]

        removed_regular = self.client.put(
            f"/api/v1/datasets/{dataset_id}/cash-position",
            headers=self.headers,
            json={
                "opening_balance": 100,
                "opening_date": "2026-09-01",
                "transactions": [],
            },
        )
        self.assertEqual(removed_regular.status_code, 403, removed_regular.text)

        changed_payload = dict(canonical_payload)
        changed_payload["transactions"] = [
            {**canonical_payload["transactions"][0], "amount": 999}
        ]
        changed_generic = self.client.put(
            f"/api/v1/datasets/{dataset_id}/records/global/cash_position",
            headers=self.headers,
            json={"payload": changed_payload},
        )
        self.assertEqual(changed_generic.status_code, 403, changed_generic.text)

        with_manual = dict(canonical_payload)
        with_manual["transactions"] = [
            *canonical_payload["transactions"],
            {
                "id": "manual-force", "date": "2026-09-09",
                "type": "余额强制修改", "amount": 88,
                "system_locked": False,
                "created_at": "2026-09-09T10:00:00+00:00",
            },
        ]
        added = self.client.put(
            f"/api/v1/datasets/{dataset_id}/records/global/cash_position",
            headers=self.headers,
            json={"payload": with_manual},
        )
        self.assertEqual(added.status_code, 200, added.text)
        without_manual = dict(added.json()["payload"])
        without_manual["transactions"] = [
            item for item in without_manual["transactions"] if item["id"] != "manual-force"
        ]
        deleted_manual = self.client.put(
            f"/api/v1/datasets/{dataset_id}/records/global/cash_position",
            headers=self.headers,
            json={"payload": without_manual},
        )
        self.assertEqual(deleted_manual.status_code, 200, deleted_manual.text)

        deleted_record = self.client.post(
            f"/api/v1/datasets/{dataset_id}/records/global/cash_position/delete",
            headers=self.headers,
            json={},
        )
        self.assertEqual(deleted_record.status_code, 403, deleted_record.text)

    def test_preopening_loan_is_reported_but_not_applied_to_balance_again(self):
        dataset_id = "sep10-preopening-loan-history"
        self.create_dataset(dataset_id, "Sep10 借款历史")
        with patch.object(
            main, "utcnow", return_value=datetime(2026, 9, 10, 4, 0, tzinfo=timezone.utc)
        ):
            saved = self.client.put(
                f"/api/v1/datasets/{dataset_id}/cash-position",
                headers=self.headers,
                json={
                    "opening_balance": 500,
                    "opening_date": "2026-09-05",
                    "transactions": [
                        {
                            "id": "old-loan", "date": "2026-08-20",
                            "type": "向股东借款", "amount": 300,
                        },
                        {
                            "id": "new-loan", "date": "2026-09-06",
                            "type": "向股东借款", "amount": 40,
                        },
                    ],
                },
            )
        self.assertEqual(saved.status_code, 200, saved.text)
        body = saved.json()
        self.assertEqual(body["current_balance"], money_text(540))
        self.assertEqual(body["shareholder_loan"], money_text(340))
        self.assertEqual(body["shareholder_loan_this_month"], money_text(40))
        self.assertEqual(body["shareholder_loan_this_year"], money_text(340))
        self.assertEqual(body["shareholder_loan_total"], money_text(340))
        self.assertNotIn("old-loan", {row.get("id") for row in body["entries"]})

    def test_audit_summary_names_added_detail_and_income_change(self):
        day = "2026-09-04"
        before = {"daily": [{"日期": day, "A店现金": 100}]}
        after = {
            "daily": [{
                "日期": day,
                "A店现金": 125,
                "食材明细": "牛肉 5斤",
            }]
        }
        summary = main._audit_change_summary(before, after, day)
        self.assertIn("增加 食材 牛肉 5斤", summary)
        self.assertIn("修改 A店现金 100.00 → 125.00", summary)

    def test_audit_summary_describes_other_record_fields(self):
        before = {"dishes": [{"日期": "2026-09-05", "菜品数据": "旧配方", "总成本": 20}]}
        after = {"dishes": [{"日期": "2026-09-05", "菜品数据": "新配方", "总成本": 25}]}
        summary = main._audit_change_summary(before, after, "2026-09-05")
        self.assertNotIn("账本内容已更新", summary)
        self.assertIn("旧配方", summary)
        self.assertIn("新配方", summary)

    def test_oct29_customers_save_after_food_or_income_and_reload(self):
        dataset_id = "oct29-customer-save"
        self.create_dataset(dataset_id, "客户数保存")
        headers = self.create_role_headers(dataset_id, "oct29-customer-save-user", "entry")
        day = "2026-09-02"
        url = f"/api/v1/datasets/{dataset_id}/records/day/{day}"
        row = {"日期": day, "A店现金": 100, "食材成本": 40, "食材明细": "牛肉 2斤 40", "新客人数量": "", "总客人数": ""}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 2, 0, tzinfo=timezone.utc)):
            saved = self.client.put(url, headers=headers, json={"payload": {"daily": [row]}})
        self.assertEqual(saved.status_code, 200, saved.text)
        with database.SessionLocal() as db:
            stored = main.get_record(db, dataset_id, "day", day)
            stored.updated_at = datetime(2026, 9, 2, 2, 0, tzinfo=timezone.utc)
            db.commit()
        # Customer counts start their own window after other fields are locked.
        row = {**row, "新客人数量": 2, "总客人数": 5}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 3, 0, tzinfo=timezone.utc)):
            saved = self.client.put(url, headers=headers, json={"payload": {"daily": [row]}})
        self.assertEqual(saved.status_code, 200, saved.text)
        reloaded = self.client.get(url, headers=headers).json()["payload"]["daily"][0]
        self.assertEqual((reloaded["新客人数量"], reloaded["总客人数"]), (2, 5))
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 3, 20, tzinfo=timezone.utc)):
            saved = self.client.put(url, headers=headers, json={"payload": {"daily": [{**row, "新客人数量": 3, "总客人数": 6}]}})
        self.assertEqual(saved.status_code, 200, saved.text)
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 3, 30, tzinfo=timezone.utc)):
            denied = self.client.put(url, headers=headers, json={"payload": {"daily": [{**row, "新客人数量": 4, "总客人数": 7}]}})
            window = self.client.get(f"/api/v1/datasets/{dataset_id}/entry-windows", headers=headers).json()[0]
        self.assertEqual(denied.status_code, 403, denied.text)
        self.assertEqual(set(window["customer_fields"]), {"new", "total"})
        self.assertTrue(all(item["state"] == "locked" for item in window["customer_fields"].values()))

    def test_oct29_customer_zero_compatibility_and_independent_first_saves(self):
        dataset_id = "oct29-customer-zero"
        self.create_dataset(dataset_id, "旧默认零兼容")
        headers = self.create_role_headers(dataset_id, "oct29-customer-zero-user", "entry")
        day = "2026-09-01"
        row = {"日期": day, "新客人数量": 0, "总客人数": 0}
        self.put_record(dataset_id, "day", day, {"daily": [row]})
        # Model a pre-upgrade row with automatic zeros and no customer clocks.
        with database.SessionLocal() as db:
            stored = main.get_record(db, dataset_id, "day", day)
            stored.updated_at = datetime(2026, 9, 1, tzinfo=timezone.utc)
            timing = main._timing_row(db, dataset_id, "day", day)
            if timing:
                db.delete(timing)
            db.commit()
        url = f"/api/v1/datasets/{dataset_id}/records/day/{day}"
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 3, 2, 0, tzinfo=timezone.utc)):
            saved = self.client.put(url, headers=headers, json={"payload": {"daily": [{**row, "总客人数": 5}]}})
        self.assertEqual(saved.status_code, 200, saved.text)
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 3, 3, 0, tzinfo=timezone.utc)):
            saved = self.client.put(url, headers=headers, json={"payload": {"daily": [{**row, "新客人数量": 2, "总客人数": 5}]}})
        self.assertEqual(saved.status_code, 200, saved.text)
        # New blank -> 0 is an intentional save and locks after 30 minutes.
        day2 = "2026-09-04"
        url2 = f"/api/v1/datasets/{dataset_id}/records/day/{day2}"
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 4, 2, 0, tzinfo=timezone.utc)):
            saved = self.client.put(url2, headers=headers, json={"payload": {"daily": [{"日期": day2, "新客人数量": 0, "总客人数": ""}]}})
        self.assertEqual(saved.status_code, 200, saved.text)
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 4, 3, 0, tzinfo=timezone.utc)):
            total_saved = self.client.put(url2, headers=headers, json={"payload": {"daily": [{"日期": day2, "新客人数量": 0, "总客人数": 10}]}})
            denied = self.client.put(url2, headers=headers, json={"payload": {"daily": [{"日期": day2, "新客人数量": 1, "总客人数": 10}]}})
        self.assertEqual(total_saved.status_code, 200, total_saved.text)
        self.assertEqual(denied.status_code, 403, denied.text)

    def test_oct29_customer_validation_and_unrelated_saves_keep_clocks(self):
        dataset_id = "oct29-customer-validation"
        self.create_dataset(dataset_id, "客户数校验")
        headers = self.create_role_headers(dataset_id, "oct29-customer-validation-user", "entry")
        day = "2026-09-02"
        url = f"/api/v1/datasets/{dataset_id}/records/day/{day}"
        for row in ({"新客人数量": -1}, {"新客人数量": 1.2}, {"总客人数": True}, {"新客人数量": 6, "总客人数": 5}):
            invalid = self.client.put(url, headers=self.headers, json={"payload": {"daily": [{"日期": day, **row}]}})
            self.assertEqual(invalid.status_code, 422, invalid.text)
        # Partial input is valid; total is independently saved later.
        row = {"日期": day, "新客人数量": 2, "总客人数": ""}
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 2, 0, tzinfo=timezone.utc)):
            saved = self.client.put(url, headers=self.headers, json={"payload": {"daily": [row]}})
        self.assertEqual(saved.status_code, 200, saved.text)
        with patch.object(main, "utcnow", return_value=datetime(2026, 9, 2, 3, 0, tzinfo=timezone.utc)):
            note_saved = self.client.put(url, headers=self.headers, json={"payload": {"daily": [{**row, "当日备注": "备注编辑"}]}})
            denied = self.client.put(url, headers=headers, json={"payload": {"daily": [{**row, "新客人数量": 3}]}})
        self.assertEqual(note_saved.status_code, 200, note_saved.text)
        self.assertEqual(denied.status_code, 403, denied.text)

    def test_oct29_food_purchase_cash_is_not_duplicated_by_inventory(self):
        dataset_id = "oct29-food-cash"
        self.create_dataset(dataset_id, "采购自动入库只扣一次现金")
        day = "2026-09-02"
        # Procurement batches are projections of this day, never second copies
        # in inventory.entries. Explicit manual receipts remain separate costs.
        self.put_record(dataset_id, "global", "inventory", {"entries": [{
            "id": "manual", "date": day, "category": "食材", "name": "土豆",
            "quantity": 2, "unit": "斤", "unit_price": 10, "total": 20,
        }]})
        for cost in (80, 80, 90, 0):
            self.put_record(dataset_id, "day", day, {"daily": [{
                "日期": day, "食材成本": cost,
                "食材明细": f"牛肉 4斤 {cost}" if cost else "",
            }]})
            with database.SessionLocal() as db:
                summary = main._cash_flow_summary_db(db, dataset_id, {
                    "opening_balance": 1000, "opening_date": "2026-09-01",
                }, now=datetime(2026, 9, 3, tzinfo=timezone.utc))
                inventory = main.get_record(db, dataset_id, "global", "inventory")
                self.assertEqual(len(inventory.payload["entries"]), 1)
            self.assertEqual(summary["inventory_expenses"], money_text(20))
            self.assertEqual(summary["daily_expenses"], money_text(cost))
            self.assertEqual(summary["current_balance"], money_text(1000 - cost - 20))


if __name__ == "__main__":
    unittest.main()
