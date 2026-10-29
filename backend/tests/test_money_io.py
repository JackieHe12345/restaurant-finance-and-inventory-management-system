import os
os.environ.setdefault("DATABASE_URL", "sqlite+pysqlite://")
os.environ.setdefault("SECRET_KEY", "dataset-api-test-secret")
os.environ.setdefault("ADMIN_PASSWORD", "Admin123!")

from decimal import Decimal
from io import BytesIO
import json
import re
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import load_workbook

from app.excel_io import (
    WorkbookFormatError, build_internal_workbook, build_user_workbook,
    merge_user_record_payload, parse_internal_workbook, parse_user_workbook,
)
from app.wechat_pay import WeChatPayError, summarize_rows, parse_trade_bill
from app.wechat_sync import recalc_revenue_breakdown

DATASET = '11111111-1111-4111-8111-111111111111'


def record(payload):
    return {'record_type': 'day', 'record_key': '2026-10-29', 'payload': payload}


def user_book(payload):
    return build_user_workbook(dataset_id=DATASET, schema_version=3, year=2026, records=[record(payload)])


def daily_result(content):
    parsed = parse_user_workbook(content, year=2026)
    return next(item.payload['daily'][0] for item in parsed.records if item.record_type == 'day')


def test_user_excel_currency_cells_are_text_and_round_trip_exact():
    content = user_book({'daily': [{'日期': '2026-10-29', '收益': '999999999999.99',
                                   '食材成本': '1.005', '新客人数量': 2,
                                   '总客人数': 3, '当日备注': '=1+1', '扩展编号': 42}]})
    book = load_workbook(BytesIO(content))
    assert book['daily']['B2'].data_type == 's'
    assert book['daily']['B2'].value == '999999999999.99'
    assert book['daily']['E2'].value == '1.01'
    book.close()
    row = daily_result(content)
    assert row['收益'] == '999999999999.99'
    assert row['食材成本'] == '1.01'
    assert row['新客人数量'] == 2
    assert row['总客人数'] == 3
    assert row['扩展编号'] == 42
    assert row['当日备注'] == '=1+1'


@pytest.mark.parametrize(('token', 'expected'), [
    ('1.004999999999999999999', '1.00'),
    ('1.005', '1.01'),
    ('999999999999.984999999', '999999999999.98'),
    ('999999999999.985', '999999999999.99'),
])
def test_numeric_excel_import_reads_original_xml_decimal_token(token, expected):
    content = user_book({'daily': [{'日期': '2026-10-29', '收益': '1.00'}]})
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                text = data.decode()
                text, count = re.subn(r'<c r="B2"[^>]*>.*?</c>', f'<c r="B2" t="n"><v>{token}</v></c>', text)
                assert count == 1
                data = text.encode()
            target.writestr(info, data)
    assert daily_result(output.getvalue())['收益'] == expected


def test_excel_nested_json_currency_and_unknown_metadata_round_trip():
    details = [{'name': '牛肉', 'quantity': 3, 'unit': '斤', 'total': '1.005',
                'unit_price': '0.3333335', 'extension': {'total': 2.718}}]
    content = user_book({'daily': [{'日期': '2026-10-29', '食材明细': json.dumps(details), '食材成本': '1.005'}]})
    row = daily_result(content)
    detail = json.loads(row['食材明细'])[0]
    assert detail['total'] == '1.01'
    assert detail['unit_price'] == '0.333334'
    assert detail['quantity'] == 3
    assert detail['extension'] == {'total': 2.718}


def test_internal_excel_preserves_export_and_normalizes_imported_legacy_currency():
    payload = {'daily': [{'日期': '2026-10-29', '收益': '999999999999.99', '食材成本': 0.1,
                         '食材明细': '[{"name":"米","total":"0.30","quantity":3}]'}],
               'unknown': {'amount': 9.87654321}}
    content = build_internal_workbook(dataset_id=DATASET, schema_version=3, records=[record(payload)])
    book = load_workbook(BytesIO(content))
    assert json.loads(book['__day_records']['C2'].value) == payload
    book.close()
    parsed = parse_internal_workbook(content)
    expected = {**payload, 'daily': [{**payload['daily'][0], '食材成本': '0.10'}]}
    assert parsed.records[0].payload == expected


def test_excel_merge_uses_decimal_for_amounts_and_quantity_price_products():
    result = merge_user_record_payload(
        {'daily': [{'日期': '2026-10-29', '收益': '0.10', '食材明细': '米 3斤 1.00'}]},
        {'daily': [{'日期': '2026-10-29', '收益': '0.20', '食材明细': '米 3斤 1.00'}]}, record_type='day')
    assert result['daily'][0]['收益'] == '0.30'
    assert result['daily'][0]['食材明细'] == '米 6斤 2.00'


@pytest.mark.parametrize('value', ['NaN', 'Infinity', '-Infinity', '1000000000000.00', True])
def test_user_excel_rejects_invalid_currency(value):
    with pytest.raises(WorkbookFormatError):
        user_book({'daily': [{'日期': '2026-10-29', '收益': value}]})


def test_wechat_summary_and_revenue_use_cents_without_float():
    summary = summarize_rows([{'订单金额': '0.1'}, {'订单金额': '0.2'}, {'订单金额': '1.005', '申请退款金额': '0.10'}])
    assert summary['amount'] == '1.21'
    assert summary['payment_amount'] == '1.31'
    assert summary['refund_amount'] == '0.10'
    row = {'A店现金': '0.10', 'A店微信自动': '0.20', 'B店现金': '1.005'}
    recalc_revenue_breakdown(row)
    assert row['A店收入'] == '0.30'
    assert row['B店收入'] == '1.01'
    assert row['收益'] == '1.31'


@pytest.mark.parametrize('value', ['NaN', 'Infinity', '-Infinity', '1000000000000.00', True])
def test_wechat_invalid_amounts_fail_closed(value):
    with pytest.raises(WeChatPayError):
        summarize_rows([{'订单金额': value}])
    if value != '1000000000000.00':
        with pytest.raises(WeChatPayError):
            recalc_revenue_breakdown({'A店现金': value})


def test_wechat_large_sum_is_exact_and_aggregate_may_exceed_single_input_cap():
    assert summarize_rows([{'订单金额': '999999999999.98'}, {'订单金额': '0.01'}])['amount'] == '999999999999.99'
    assert summarize_rows([{'订单金额': '999999999999.99'}, {'订单金额': '0.01'}])['amount'] == '1000000000000.00'
    row = {'A店微信自动': '1000000000000.00', 'B店现金': '0.10'}
    recalc_revenue_breakdown(row)
    assert row['收益'] == '1000000000000.10'
    assert row['A店收入'] == '1000000000000.00'
    csv = '\n'.join([
        '交易时间,微信订单号,商户号,订单金额,申请退款金额',
        '2026-10-28 10:00:00,1,m,999999999999.99,0',
        '2026-10-28 11:00:00,2,m,0.01,0',
    ])
    bill = parse_trade_bill(csv, '2026-10-28')
    assert bill.totals.order_amount == Decimal('1000000000000.00')
    assert bill.totals.order_net == Decimal('1000000000000.00')


def test_excel_merge_json_details_preserves_purchase_batches_and_unknown_metadata():
    before = '[{"name":"米","quantity":3,"unit":"斤","total":1.004999999999999999,"extension":{"tag":"keep"}}]'
    result = merge_user_record_payload(
        {'daily': [{'日期': '2026-10-29', '收益': '0.10', '食材明细': before}]},
        {'daily': [{'日期': '2026-10-29', '收益': '0.20', '食材明细': [{'name': '米', 'quantity': 2, 'unit': '斤', 'total': '2.30'}]}]}, record_type='day')
    rows = json.loads(result['daily'][0]['食材明细'])
    assert rows == [{'name': '米', 'quantity': 3, 'unit': '斤', 'total': '1.00', 'extension': {'tag': 'keep'}},
                    {'name': '米', 'quantity': 2, 'unit': '斤', 'total': '2.30'}]


def test_excel_mixed_json_text_merge_keeps_all_purchase_quantities_and_totals():
    result = merge_user_record_payload(
        {'daily': [{'日期': '2026-10-29', '食材明细': '米 3斤 1.005'}]},
        {'daily': [{'日期': '2026-10-29', '食材明细': '[{"name":"米","quantity":2,"unit":"斤","total":"0.20"}]'}]}, record_type='day')
    rows = json.loads(result['daily'][0]['食材明细'])
    assert [row['quantity'] for row in rows] == [3, 2]
    assert [row['total'] for row in rows] == ['1.01', '0.20']
    assert rows[0]['legacy_text'] == '米 3斤 1.005'


def test_internal_excel_numeric_json_currency_is_read_before_float_conversion():
    content = build_internal_workbook(dataset_id=DATASET, schema_version=3,
                                     records=[record({'daily': [{'日期': '2026-10-29', '收益': '1.00'}]})])
    book = load_workbook(BytesIO(content))
    book['__day_records']['C2'] = '{"daily":[{"日期":"2026-10-29","收益":1.004999999999999999,"食材明细":[{"name":"米","quantity":1.5,"total":999999999999.98499999,"unit_price":1.0000004999999}]}],"extension":{"amount":1.234}}'
    output = BytesIO()
    book.save(output)
    book.close()
    payload = parse_internal_workbook(output.getvalue()).records[0].payload
    row = payload['daily'][0]
    assert row['收益'] == '1.00'
    assert row['食材明细'][0]['total'] == '999999999999.98'
    assert row['食材明细'][0]['unit_price'] == '1.000000'
    assert row['食材明细'][0]['quantity'] == 1.5
    assert payload['extension'] == {'amount': 1.234}
