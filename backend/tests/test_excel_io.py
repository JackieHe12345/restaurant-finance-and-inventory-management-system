from __future__ import annotations

import json
import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.excel_io import (
    MAX_JSON_CELL_CHARS,
    USER_SHEETS,
    WorkbookFormatError,
    WorkbookRecord,
    build_internal_workbook,
    build_user_workbook,
    canonical_json,
    parse_internal_workbook,
    parse_user_workbook,
    payload_checksum,
)


DATASET_ID = "11111111-1111-4111-8111-111111111111"


def record(record_type: str, key: str, payload: dict) -> WorkbookRecord:
    return WorkbookRecord(record_type, key, payload, payload_checksum(payload))


def user_workbook_with_rows(rows_by_sheet: dict[str, list[list]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, headers in USER_SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(list(headers))
        for row in rows_by_sheet.get(sheet_name, []):
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class ExcelIoTests(unittest.TestCase):
    def test_empty_internal_workbook_is_valid(self):
        content = build_internal_workbook(
            dataset_id=DATASET_ID, schema_version=2, records=[]
        )
        parsed = parse_internal_workbook(content)
        self.assertEqual(parsed.records, [])
        self.assertEqual(parsed.source_dataset_id, DATASET_ID)

    def test_audit_and_cash_checkbox_are_internal_only(self):
        records = [
            record(
                "day",
                "2026-08-01",
                {
                    "daily": [{"日期": "2026-08-01", "收益": 10}],
                    "fixed_expenses": [
                        {
                            "日期": "2026-08-01",
                            "类别": "房租",
                            "金额": 100,
                            # Simulate a stale cached web alias.  The App's
                            # explicit Chinese False must be authoritative.
                            "include_cash_balance": True,
                            "计入可用余额": False,
                        }
                    ],
                },
            )
        ]
        audit = {
            "id": "audit-1",
            "dataset_id": DATASET_ID,
            "created_at": "2026-09-06T10:00:00+08:00",
            "actor_username": "admin",
            "change_summary": "修改固定成本",
        }
        internal = build_internal_workbook(
            dataset_id=DATASET_ID,
            schema_version=2,
            records=records,
            audit_records=[audit],
        )
        workbook = load_workbook(BytesIO(internal), data_only=True)
        self.assertIn("__audit_records", workbook.sheetnames)
        self.assertEqual(workbook["__audit_records"].sheet_state, "veryHidden")
        self.assertEqual(workbook["__audit_records"]["B2"].value, "audit-1")
        saved_payload = json.loads(workbook["__day_records"]["C2"].value)
        saved_expense = saved_payload["fixed_expenses"][0]
        self.assertIs(saved_expense["计入可用余额"], False)
        self.assertNotIn("include_cash_balance", saved_expense)
        workbook.close()

        # Optional audit cache does not make an internal backup incompatible.
        parsed_internal = parse_internal_workbook(internal)
        self.assertEqual(parsed_internal.records[0].payload["fixed_expenses"][0]["计入可用余额"], False)

        user = build_user_workbook(
            dataset_id=DATASET_ID, schema_version=2, year=2026, records=records
        )
        user_workbook = load_workbook(BytesIO(user), data_only=True)
        self.assertNotIn("__audit_records", user_workbook.sheetnames)
        user_headers = [cell.value for cell in user_workbook["fixed_expenses"][1]]
        self.assertNotIn("计入可用余额", user_headers)
        self.assertNotIn("include_cash_balance", user_headers)
        user_workbook.close()

        parsed_user = parse_user_workbook(user, year=2026)
        expense = parsed_user.records[0].payload["fixed_expenses"][0]
        self.assertNotIn("计入可用余额", expense)
        self.assertNotIn("include_cash_balance", expense)

    def test_internal_round_trip_keeps_arbitrary_global_and_recalculates_bad_sha(self):
        source = [
            record("day", "2026-08-01", {"daily": [{"日期": "2026-08-01", "收益": 10}]}),
            record("month", "2026-08", {"monthly_fixed": []}),
            record("global", "audit/custom", {"中文": "保留", "rows": [1, 2]}),
        ]
        content = build_internal_workbook(
            dataset_id=DATASET_ID, schema_version=2, records=source
        )
        workbook = load_workbook(BytesIO(content))
        workbook["__day_records"]["D2"] = "0" * 64
        changed = BytesIO()
        workbook.save(changed)
        workbook.close()

        parsed = parse_internal_workbook(changed.getvalue())
        self.assertEqual(
            {(item.record_type, item.record_key) for item in parsed.records},
            {("day", "2026-08-01"), ("month", "2026-08"), ("global", "audit/custom")},
        )
        self.assertTrue(parsed.warnings)

    def test_internal_requires_all_four_hidden_tables(self):
        content = build_internal_workbook(
            dataset_id=DATASET_ID, schema_version=2, records=[]
        )
        workbook = load_workbook(BytesIO(content))
        workbook.remove(workbook["__day_records"])
        changed = BytesIO()
        workbook.save(changed)
        workbook.close()
        with self.assertRaises(WorkbookFormatError):
            parse_internal_workbook(changed.getvalue())

    def test_json_cell_limit_is_not_silently_truncated(self):
        overhead = len(canonical_json({"x": ""}))
        exact = {"x": "a" * (MAX_JSON_CELL_CHARS - overhead)}
        content = build_internal_workbook(
            dataset_id=DATASET_ID,
            schema_version=2,
            records=[record("global", "boundary", exact)],
        )
        self.assertTrue(parse_internal_workbook(content).records)
        too_large = {"x": "a" * (MAX_JSON_CELL_CHARS - overhead + 1)}
        with self.assertRaises(WorkbookFormatError):
            build_internal_workbook(
                dataset_id=DATASET_ID,
                schema_version=2,
                records=[record("global", "too-large", too_large)],
            )

    def test_user_round_trip_and_formula_like_text_stays_text(self):
        records = [
            record(
                "day",
                "2026-08-01",
                {
                    "daily": [{"日期": "wrong", "收益": 10, "当日备注": "=1+1"}],
                    "fixed_expenses": [{"日期": "wrong", "类别": "水费", "金额": 2}],
                    "dishes": [],
                },
            ),
            record("month", "2026-08", {"monthly_fixed": [{"月份": "wrong", "人工总额": 3}]}),
            record("global", "ingredients", {"ingredients": [{"词条名称": "米", "分类": "食材"}]}),
            record("global", "fixed_categories", {"fixed_categories": []}),
        ]
        content = build_user_workbook(
            dataset_id=DATASET_ID,
            schema_version=2,
            year=2026,
            records=records,
        )
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(workbook["daily"]["L2"].value, "=1+1")
        self.assertEqual(workbook["daily"]["L2"].data_type, "s")
        self.assertEqual(workbook["__app_meta"].sheet_state, "veryHidden")
        workbook.close()

        parsed = parse_user_workbook(content, year=2026)
        by_key = {(item.record_type, item.record_key): item for item in parsed.records}
        self.assertEqual(by_key[("day", "2026-08-01")].payload["daily"][0]["日期"], "2026-08-01")
        self.assertIn(("global", "fixed_categories"), by_key)
        with self.assertRaises(WorkbookFormatError):
            parse_user_workbook(content, year=2025)

    def test_user_replacement_requires_all_six_sheets(self):
        workbook = Workbook()
        daily = workbook.active
        daily.title = "daily"
        daily.append(["日期", "收益"])
        daily.append(["2026-08-01", 12])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        with self.assertRaisesRegex(WorkbookFormatError, "完整六张工作表"):
            parse_user_workbook(output.getvalue(), year=2026)

    def test_user_replacement_requires_complete_standard_columns(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, headers in USER_SHEETS.items():
            sheet = workbook.create_sheet(sheet_name)
            sheet.append([headers[0]])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        with self.assertRaisesRegex(WorkbookFormatError, "缺少列"):
            parse_user_workbook(output.getvalue(), year=2026)

    def test_user_mixed_years_is_rejected_for_authoritative_replacement(self):
        content = user_workbook_with_rows(
            {
                "daily": [
                    ["2025-12-31", 5],
                    ["2026-01-01", 6],
                ]
            }
        )
        with self.assertRaisesRegex(WorkbookFormatError, "只能包含 2026 年"):
            parse_user_workbook(content, year=2026)

    def test_user_family_id_must_match_target_when_present(self):
        content = build_user_workbook(
            dataset_id=DATASET_ID,
            schema_version=2,
            year=2026,
            records=[],
        )
        parsed = parse_user_workbook(
            content, year=2026, expected_dataset_id=DATASET_ID
        )
        self.assertEqual(parsed.source_dataset_id, DATASET_ID)
        with self.assertRaisesRegex(WorkbookFormatError, "不能替换当前账本"):
            parse_user_workbook(
                content, year=2026, expected_dataset_id="another-ledger"
            )

    def test_user_formula_cells_are_rejected_but_literal_equals_text_is_allowed(self):
        content = user_workbook_with_rows({"daily": [["2026-01-01", "=1+1"]]})
        workbook = load_workbook(BytesIO(content))
        workbook["daily"]["B2"] = "=1+1"
        formula_output = BytesIO()
        workbook.save(formula_output)
        workbook.close()
        with self.assertRaisesRegex(WorkbookFormatError, "含有公式"):
            parse_user_workbook(formula_output.getvalue(), year=2026)

        workbook = load_workbook(BytesIO(content))
        workbook["daily"]["B2"] = "'=1+1"
        text_output = BytesIO()
        workbook.save(text_output)
        workbook.close()
        with self.assertRaisesRegex(WorkbookFormatError, "金额无效"):
            parse_user_workbook(text_output.getvalue(), year=2026)
        # Formula-like text remains valid in a note, never as a monetary value.
        workbook = load_workbook(BytesIO(content))
        workbook["daily"]["B2"] = "2.00"
        workbook["daily"]["L2"] = "'=1+1"
        note_output = BytesIO()
        workbook.save(note_output)
        workbook.close()
        self.assertTrue(parse_user_workbook(note_output.getvalue(), year=2026).records)


if __name__ == "__main__":
    unittest.main()
