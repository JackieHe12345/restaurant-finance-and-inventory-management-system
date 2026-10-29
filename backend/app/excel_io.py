from __future__ import annotations

import hashlib
import json
import math
import re
import posixpath
import xml.etree.ElementTree as ET
from decimal import Decimal, InvalidOperation
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import date, datetime, time, timezone
from io import BytesIO
from typing import Any, Iterable
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .money import money, money_text, normalize_payload, exact_request_numbers, storage_json_dumps, SCHEMA_FIELDS


MAX_ZIP_MEMBERS = 2_000
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_RECORDS = 50_000
MAX_JSON_CELL_CHARS = 32_767
MAX_CELL_CHARS = 32_767
MAX_RECORD_JSON_BYTES = 2 * 1024 * 1024

INTERNAL_SHEETS = {
    "__day_records": ("dataset_id", "日期", "payload_json", "sha256"),
    "__month_records": ("dataset_id", "月份", "payload_json", "sha256"),
    "__global_records": ("dataset_id", "记录类型", "payload_json", "sha256"),
    "__meta": ("dataset_id", "schema_version", "last_synced_at"),
}
INTERNAL_AUDIT_SHEET = "__audit_records"
INTERNAL_AUDIT_COLUMNS = (
    "dataset_id", "audit_id", "created_at", "payload_json", "sha256"
)

# This checkbox is business metadata, rather than user-Excel data.  Older web
# builds wrote the English alias; accept it on input, but persist only the
# App's canonical Chinese field.
FIXED_EXPENSE_CASH_FIELD = "计入可用余额"
FIXED_EXPENSE_CASH_LEGACY_FIELD = "include_cash_balance"
INTERNAL_FIXED_EXPENSE_FIELDS = frozenset(
    {FIXED_EXPENSE_CASH_FIELD, FIXED_EXPENSE_CASH_LEGACY_FIELD}
)

USER_SHEETS: dict[str, tuple[str, ...]] = {
    "daily": (
        # Keep the legacy 12 columns first so old user Excel layouts remain stable.
        "日期", "收益", "新客人数量", "总客人数", "食材成本", "食材明细",
        "佐料成本", "佐料明细", "水电", "其他", "其他明细", "当日备注",
        # v24/v25 store and payment breakdown fields are appended.
        "总现金", "A店现金", "B店现金", "总线上",
        "A店微信增加", "A店微信自动", "A店微信覆盖", "A店支付宝增加", "A店支付宝覆盖",
        "B店微信增加", "B店微信自动", "B店微信覆盖", "B店支付宝增加", "B店支付宝覆盖",
        "A店收入", "B店收入", "学生外卖", "A店外卖", "B店外卖", "手续费",
        "A店微信支付金额", "A店微信退款金额", "A店微信支付笔数", "A店微信退款笔数",
        "B店微信支付金额", "B店微信退款金额", "B店微信支付笔数", "B店微信退款笔数",
        "微信账单日期", "微信账单同步时间", "微信账单口径", "微信账单SHA1", "微信账单未归属笔数", "微信账单强制更新",
    ),
    "monthly_fixed": (
        # Keep the legacy fixed-cost columns in their original order.
        "月份", "人工总额", "房租总额", "水费支出", "电费支出", "水电费",
        "实际使用水费", "实际使用电费", "煤气费", "其它固定成本", "月度备注",
        "预计水费", "预计电费", "预计煤气费",
    ),
    # Whether a fixed expense affects the cash balance is intentionally not in
    # a user workbook.  It is synchronised in the internal workbook only.
    "fixed_expenses": ("日期", "类别", "明细", "金额"),
    "fixed_categories": ("类别名称",),
    "ingredients": ("词条名称", "分类"),
    "dishes": ("日期", "菜品明细", "菜品数据", "总成本", "备注"),
}

# Exact schema fields, never a heuristic that turns ids, quantities or custom
# metadata into currency. User workbooks store currency as decimal text.
_USER_MONEY_FIELDS = set().union(*(set(SCHEMA_FIELDS.get(name, ())) for name in USER_SHEETS))


DATE_SHEETS = {"daily": "日期", "fixed_expenses": "日期", "dishes": "日期"}
DAY_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")
MONTH_PATTERN = re.compile(r"^\d{4}-\d{2}$")
CONTROL_PATTERN = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class WorkbookFormatError(ValueError):
    pass


@dataclass(frozen=True)
class WorkbookRecord:
    record_type: str
    record_key: str
    payload: dict[str, Any]
    checksum: str


@dataclass
class ParsedWorkbook:
    records: list[WorkbookRecord]
    warnings: list[str] = field(default_factory=list)
    source_dataset_id: str | None = None
    present_sheets: set[str] = field(default_factory=set)


_PRICE_FIELD_RE = re.compile(r"(?:^|[_-])(unit_price|price)(?:$|[_-])|单价|单位价")
_DETAIL_FIELDS = {"食材明细", "佐料明细", "其他明细"}


def _number(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    try:
        number = Decimal(str(value).strip())
    except (TypeError, ValueError, InvalidOperation):
        return None
    return number if number.is_finite() else None


def fixed_expense_counts_towards_cash(expense: dict[str, Any]) -> bool:
    """Resolve the internal cash-balance flag without letting stale aliases win.

    The desktop App uses ``计入可用余额``.  A short-lived web build wrote
    ``include_cash_balance`` instead, and some synced rows can contain both.
    Presence matters here: ``False`` is a valid, explicit App choice and must
    never fall through to a stale English ``True``.  Rows created before the
    checkbox existed retain the historic default of being included.
    """

    if not isinstance(expense, dict):
        return True
    if (
        FIXED_EXPENSE_CASH_FIELD in expense
        and expense.get(FIXED_EXPENSE_CASH_FIELD) not in (None, "")
    ):
        value = expense.get(FIXED_EXPENSE_CASH_FIELD)
    elif (
        FIXED_EXPENSE_CASH_LEGACY_FIELD in expense
        and expense.get(FIXED_EXPENSE_CASH_LEGACY_FIELD) not in (None, "")
    ):
        value = expense.get(FIXED_EXPENSE_CASH_LEGACY_FIELD)
    else:
        return True

    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return bool(value)
    normalized = str(value).strip().casefold()
    if normalized in {"false", "0", "no", "n", "否", "不", "off"}:
        return False
    if normalized in {"true", "1", "yes", "y", "是", "on"}:
        return True
    # Unknown legacy values should retain the original behaviour (included),
    # rather than silently hiding a cost from the cash-balance table.
    return True


def normalize_fixed_expense_cash_flags(payload: dict[str, Any]) -> dict[str, Any]:
    """Return a payload with one resolved cash-inclusion value per expense.

    This gives all sync paths one deterministic internal representation and
    ensures a canonical Chinese ``False`` wins over an old English alias that
    says ``True``.  The legacy alias is deliberately removed so it cannot
    reappear in a future client merge or internal Excel export.
    """

    if not isinstance(payload, dict):
        return payload
    expenses = payload.get("fixed_expenses")
    if not isinstance(expenses, list):
        return deepcopy(payload)
    normalized_payload = deepcopy(payload)
    normalized_expenses: list[Any] = []
    for expense in expenses:
        if not isinstance(expense, dict):
            normalized_expenses.append(deepcopy(expense))
            continue
        normalized_expense = deepcopy(expense)
        included = fixed_expense_counts_towards_cash(normalized_expense)
        normalized_expense[FIXED_EXPENSE_CASH_FIELD] = included
        normalized_expense.pop(FIXED_EXPENSE_CASH_LEGACY_FIELD, None)
        normalized_expenses.append(normalized_expense)
    normalized_payload["fixed_expenses"] = normalized_expenses
    return normalized_payload


def _fixed_expense_match_keys(expense: dict[str, Any]) -> tuple[tuple[str, ...], tuple[str, ...]]:
    """Return exact and resilient user-Excel matching keys for one expense."""

    date_value = str(expense.get("日期", expense.get("date", "")) or "").strip()
    category = str(expense.get("类别", expense.get("category", "")) or "").strip()
    detail = str(expense.get("明细", expense.get("description", "")) or "").strip()
    amount = str(expense.get("金额", expense.get("amount", "")) or "").strip()
    return (date_value, category, detail, amount), (date_value, category, detail)


def preserve_internal_fixed_expense_flags(
    existing_payload: dict[str, Any] | None,
    incoming_payload: dict[str, Any],
) -> dict[str, Any]:
    """Keep hidden checkbox choices during an authoritative user-Excel import.

    User workbooks deliberately omit the checkbox.  An unchanged exported row
    must therefore retain its saved internal value when it is uploaded again.
    Prefer an exact business-row match and fall back to date/category/detail so
    a user may correct an amount without unexpectedly changing cash treatment.
    New or unrecognisable rows retain the legacy default (included).
    """

    incoming = deepcopy(incoming_payload)
    if not isinstance(existing_payload, dict):
        return normalize_fixed_expense_cash_flags(incoming)
    existing_expenses = existing_payload.get("fixed_expenses")
    incoming_expenses = incoming.get("fixed_expenses")
    if not isinstance(existing_expenses, list) or not isinstance(incoming_expenses, list):
        return normalize_fixed_expense_cash_flags(incoming)

    by_exact: dict[tuple[str, ...], list[dict[str, Any]]] = {}
    by_core: dict[tuple[str, ...], list[dict[str, Any]]] = {}
    for expense in existing_expenses:
        if not isinstance(expense, dict):
            continue
        resolved = deepcopy(expense)
        included = fixed_expense_counts_towards_cash(resolved)
        resolved[FIXED_EXPENSE_CASH_FIELD] = included
        resolved.pop(FIXED_EXPENSE_CASH_LEGACY_FIELD, None)
        exact, core = _fixed_expense_match_keys(resolved)
        by_exact.setdefault(exact, []).append(resolved)
        by_core.setdefault(core, []).append(resolved)

    restored: list[Any] = []
    used_ids: set[int] = set()

    def first_unused(candidates: list[dict[str, Any]] | None) -> dict[str, Any] | None:
        for candidate in candidates or []:
            candidate_id = id(candidate)
            if candidate_id not in used_ids:
                used_ids.add(candidate_id)
                return candidate
        return None

    for expense in incoming_expenses:
        if not isinstance(expense, dict):
            restored.append(deepcopy(expense))
            continue
        restored_expense = deepcopy(expense)
        exact, core = _fixed_expense_match_keys(restored_expense)
        previous = first_unused(by_exact.get(exact)) or first_unused(by_core.get(core))
        if previous is not None:
            included = fixed_expense_counts_towards_cash(previous)
            restored_expense[FIXED_EXPENSE_CASH_FIELD] = included
            restored_expense.pop(FIXED_EXPENSE_CASH_LEGACY_FIELD, None)
        restored.append(restored_expense)

    incoming["fixed_expenses"] = restored
    return normalize_fixed_expense_cash_flags(incoming)


def _display_number(value: Decimal) -> int | str:
    # Counts/quantities may remain fractional; never route their monetary
    # neighbours through a binary float while merging an Excel upload.
    return int(value) if value == value.to_integral_value() else format(value.normalize(), "f")


def _detail_item(value: str) -> tuple[str, Decimal, str, Decimal, str] | None:
    """Parse the desktop detail format using the original decimal tokens."""
    text = str(value or "").strip()
    if not text or text.lower() == "nan":
        return None
    text = re.sub(r"\s*[（(]([^）)]*)[）)]\s*$", "", text).strip()
    match = re.match(
        r"^(.*?)\s+(\d+(?:\.\d+)?)\s*([^\d\s]+)?\s+(\d+(?:\.\d+)?)$",
        text,
    )
    if match:
        name, quantity, unit, total = match.groups()
        return name.strip(), Decimal(quantity), (unit or "").strip(), money(total), text
    match = re.match(r"^(.*?)\s*(\d+(?:\.\d+)?)$", text)
    if match:
        name, total = match.groups()
        return name.strip(), Decimal("0"), "", money(total), text
    return None


def _structured_detail(raw: Any) -> tuple[list[dict[str, Any]] | None, bool]:
    if raw in (None, ""):
        return [], False
    structured = isinstance(raw, list) or isinstance(raw, str) and raw.lstrip().startswith(("[", "{"))
    if not structured:
        return None, False
    try:
        normalized = normalize_payload({"daily": [{"食材明细": raw}]})["daily"][0]["食材明细"]
        rows = json.loads(normalized) if isinstance(normalized, str) else normalized
    except (ValueError, TypeError) as exc:
        raise WorkbookFormatError("采购 JSON 明细无法识别，已取消合并。") from exc
    if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
        raise WorkbookFormatError("采购 JSON 明细必须是对象数组，已取消合并。")
    return rows, True


def _merge_detail_text(existing: Any, incoming: Any) -> Any:
    old_rows, old_structured = _structured_detail(existing)
    new_rows, new_structured = _structured_detail(incoming)
    if old_structured or new_structured:
        # JSON detail rows carry their own batch price and unknown metadata.
        # Append those purchases intact instead of splitting JSON at commas or
        # silently replacing one batch's amount with another batch's unit price.
        combined = []
        for raw, rows in ((existing, old_rows), (incoming, new_rows)):
            if rows is None:
                rows = []
                for token in re.split(r"[,，;；\n]", str(raw or "")):
                    if not token.strip():
                        continue
                    parsed = _detail_item(token)
                    if parsed is None:
                        raise WorkbookFormatError("混合采购明细存在无法识别的文本，已取消合并；请先核对明细格式。")
                    name, quantity, unit, total, _ = parsed
                    rows.append({"name": name, "quantity": _display_number(quantity), "unit": unit,
                                 "total": money_text(total), "legacy_text": token.strip()})
            combined.extend(deepcopy(rows))
        return combined if isinstance(existing, list) else json.dumps(combined, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    parts: list[str] = []
    for raw in (existing, incoming):
        text = str(raw or "").strip()
        if text and text.lower() != "nan":
            parts.extend(item.strip() for item in re.split(r"[,，]", text) if item.strip())
    merged: dict[tuple[str, str], dict[str, Any]] = {}
    fallback: list[str] = []
    for part in parts:
        parsed = _detail_item(part)
        if parsed is None:
            if part not in fallback:
                fallback.append(part)
            continue
        name, quantity, unit, total, original = parsed
        identity = (name, unit)
        item = merged.get(identity)
        if item is None:
            merged[identity] = {
                "name": name,
                "quantity": quantity,
                "unit": unit,
                "total": total,
                "has_quantity": quantity > 0,
            }
            continue
        if item["has_quantity"] and quantity > 0:
            previous_quantity = item["quantity"]
            previous_total = item["total"]
            original_unit_price = (
                previous_total / previous_quantity if previous_quantity > 0 else 0
            )
            item["quantity"] += quantity
            # Keep the original unit price.  Recalculate the combined total
            # from it instead of adding unit prices from two uploads.
            item["total"] = item["quantity"] * original_unit_price
        else:
            item["total"] += total
    output: list[str] = []
    for item in merged.values():
        if item["has_quantity"]:
            quantity = _display_number(item["quantity"])
            total = money_text(item["total"])
            unit_joined = f"{quantity}{item['unit']}" if item["unit"] else str(quantity)
            output.append(f"{item['name']} {unit_joined} {total}")
        else:
            output.append(f"{item['name']} {money_text(item['total'])}")
    output.extend(fallback)
    return ", ".join(output)


def _merge_row(existing: dict[str, Any], incoming: dict[str, Any]) -> dict[str, Any]:
    result = deepcopy(existing)
    for key, new_value in incoming.items():
        if key in _DETAIL_FIELDS:
            result[key] = _merge_detail_text(result.get(key), new_value)
            continue
        old_value = result.get(key)
        old_number, new_number = _number(old_value), _number(new_value)
        if old_number is not None and new_number is not None:
            if _PRICE_FIELD_RE.search(str(key)):
                result[key] = old_value
            else:
                result[key] = (money_text(money(old_number) + money(new_number)) if key in _USER_MONEY_FIELDS else _display_number(old_number + new_number))
        elif old_value in (None, ""):
            result[key] = deepcopy(new_value)
        elif new_value not in (None, "") and str(old_value) != str(new_value):
            if str(key) in {"当日备注", "月度备注", "备注", "明细"}:
                result[key] = f"{old_value}；{new_value}"
    return result


def merge_user_record_payload(
    existing_payload: dict[str, Any] | None,
    incoming_payload: dict[str, Any],
    *,
    record_type: str,
) -> dict[str, Any]:
    """Merge one parsed user workbook record into an existing cloud payload."""
    if not existing_payload:
        return deepcopy(incoming_payload)
    result = deepcopy(existing_payload)
    if record_type == "day":
        for field in ("daily", "monthly_fixed"):
            if field in incoming_payload:
                incoming_rows = incoming_payload.get(field) or []
                existing_rows = result.get(field) or []
                if isinstance(incoming_rows, list) and isinstance(existing_rows, list):
                    if field == "daily":
                        if existing_rows or incoming_rows:
                            combined = deepcopy(existing_rows[0]) if existing_rows else {}
                            for row in existing_rows[1:] + incoming_rows:
                                combined = _merge_row(combined, row)
                            result[field] = [combined]
                        else:
                            result[field] = deepcopy(existing_rows or incoming_rows)
        for field in ("fixed_expenses", "dishes"):
            old_rows = result.get(field) or []
            new_rows = incoming_payload.get(field) or []
            if isinstance(old_rows, list) and isinstance(new_rows, list):
                result[field] = deepcopy(old_rows) + deepcopy(new_rows)
        for field, value in incoming_payload.items():
            if field not in {"daily", "fixed_expenses", "dishes"} and field not in result:
                result[field] = deepcopy(value)
        return result
    if record_type == "month":
        old_rows = result.get("monthly_fixed") or []
        new_rows = incoming_payload.get("monthly_fixed") or []
        if old_rows or new_rows:
            combined = deepcopy(old_rows[0]) if old_rows else {}
            for row in old_rows[1:] + new_rows:
                combined = _merge_row(combined, row)
            result["monthly_fixed"] = [combined]
        else:
            result["monthly_fixed"] = deepcopy(old_rows or new_rows)
        return result
    if record_type == "global":
        for field, value in incoming_payload.items():
            old_rows = result.get(field) or []
            if isinstance(old_rows, list) and isinstance(value, list):
                seen = {canonical_json(item) for item in old_rows if isinstance(item, dict)}
                result[field] = deepcopy(old_rows) + [
                    deepcopy(item) for item in value
                    if canonical_json(item) not in seen
                ]
            elif field not in result:
                result[field] = deepcopy(value)
        return result
    return _merge_row(result, incoming_payload)


def canonical_json(payload: dict[str, Any]) -> str:
    try:
        return json.dumps(
            payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        )
    except (TypeError, ValueError, RecursionError) as exc:
        raise WorkbookFormatError(f"记录内容不能转换为有效 JSON：{exc}") from exc


def payload_checksum(payload: dict[str, Any]) -> str:
    return hashlib.sha256(canonical_json(payload).encode("utf-8")).hexdigest()


def validate_xlsx_container(content: bytes) -> None:
    if not content.startswith(b"PK\x03\x04"):
        raise WorkbookFormatError("文件不是有效的 .xlsx（ZIP）工作簿。")
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                raise WorkbookFormatError("工作簿内部文件数量异常，已拒绝读取。")
            total = 0
            for item in members:
                normalized = item.filename.replace("\\", "/")
                if normalized.startswith("/") or ".." in normalized.split("/"):
                    raise WorkbookFormatError("工作簿包含不安全的内部路径。")
                total += int(item.file_size)
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise WorkbookFormatError("工作簿解压后过大，已拒绝读取。")
                lowered = normalized.lower()
                if lowered.endswith("vbaproject.bin") or "/embeddings/" in lowered:
                    raise WorkbookFormatError("不接受包含宏或嵌入对象的工作簿。")
            names = {item.filename for item in members}
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise WorkbookFormatError("文件缺少 Excel 工作簿结构。")
    except BadZipFile as exc:
        raise WorkbookFormatError("文件不是有效的 .xlsx 工作簿。") from exc


def _numeric_xml_tokens(content: bytes) -> dict[str, dict[str, str]]:
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    result: dict[str, dict[str, str]] = {}
    with ZipFile(BytesIO(content)) as archive:
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        paths = {item.attrib.get("Id"): item.attrib.get("Target", "")
                 for item in relationships if item.attrib.get("TargetMode") != "External"}
        root = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet in root.findall("m:sheets/m:sheet", ns):
            target = paths.get(sheet.attrib.get("{" + ns["r"] + "}id"), "")
            path = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            if path not in archive.namelist():
                continue
            tokens: dict[str, str] = {}
            with archive.open(path) as source:
                for _event, cell in ET.iterparse(source, events=("end",)):
                    if cell.tag != "{" + ns["m"] + "}c":
                        continue
                    value = cell.find("m:v", ns)
                    if cell.attrib.get("t", "n") == "n" and value is not None and value.text is not None:
                        tokens[cell.attrib.get("r", "")] = value.text
                    cell.clear()
            result[sheet.attrib["name"]] = tokens
    return result


def _open_workbook(content: bytes, *, data_only: bool):
    validate_xlsx_container(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=data_only,
            keep_links=False,
        )
        # openpyxl converts numeric OOXML cells to float before callers see
        # them. Keep their original tokens for monetary columns so 1.005 or a
        # large exact decimal is rounded once, from the source text.
        for sheet_name, tokens in _numeric_xml_tokens(content).items():
            if sheet_name in workbook.sheetnames:
                workbook[sheet_name]._money_decimal_tokens = tokens
        return workbook
    except Exception as exc:
        raise WorkbookFormatError(f"无法读取 Excel 工作簿：{exc}") from exc


def _header_text(value: Any) -> str:
    return str(value or "").strip()


def _table_rows(
    sheet: Any,
    *,
    required: Iterable[str],
    aliases: dict[str, str] | None = None,
) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    iterator = sheet.iter_rows(values_only=True)
    try:
        first = next(iterator)
    except StopIteration as exc:
        raise WorkbookFormatError(f"工作表「{sheet.title}」缺少表头。") from exc
    headers = [_header_text(value) for value in first]
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise WorkbookFormatError(f"工作表「{sheet.title}」缺少表头。")

    aliases = aliases or {}
    if aliases:
        original = set(headers)
        headers = [aliases.get(item, item) if aliases.get(item, item) not in original else item for item in headers]
    nonempty_headers = [item for item in headers if item]
    if len(nonempty_headers) != len(set(nonempty_headers)):
        raise WorkbookFormatError(f"工作表「{sheet.title}」存在重复列名。")
    missing = [item for item in required if item not in headers]
    if missing:
        raise WorkbookFormatError(
            f"工作表「{sheet.title}」缺少列：{'、'.join(missing)}。"
        )

    rows: list[tuple[int, dict[str, Any]]] = []
    for row_number, values in enumerate(iterator, start=2):
        values = tuple(values)
        if not any(value not in (None, "") for value in values):
            continue
        row = {
            header: (
                getattr(sheet, "_money_decimal_tokens", {}).get(
                    f"{get_column_letter(index + 1)}{row_number}",
                    values[index] if index < len(values) else None,
                ) if header in _USER_MONEY_FIELDS and sheet.title in USER_SHEETS
                else values[index] if index < len(values) else None
            )
            for index, header in enumerate(headers)
            if header
        }
        rows.append((row_number, row))
        if len(rows) > MAX_RECORDS:
            raise WorkbookFormatError("工作簿记录数超过安全上限。")
    return headers, rows


def _valid_dataset_id(value: Any, *, location: str) -> str:
    result = str(value or "").strip()
    if not result or len(result) > 64 or CONTROL_PATTERN.search(result):
        raise WorkbookFormatError(f"{location} 的 dataset_id 无效。")
    return result


def _valid_day(value: Any, *, location: str) -> str:
    if isinstance(value, datetime):
        result = value.date().isoformat()
    elif isinstance(value, date):
        result = value.isoformat()
    else:
        result = str(value or "").strip()
    if not DAY_PATTERN.fullmatch(result):
        raise WorkbookFormatError(f"{location} 的日期无效，应为 YYYY-MM-DD。")
    try:
        if date.fromisoformat(result).isoformat() != result:
            raise ValueError
    except ValueError as exc:
        raise WorkbookFormatError(f"{location} 的日期无效：{result}。") from exc
    return result


def _valid_month(value: Any, *, location: str) -> str:
    if isinstance(value, (datetime, date)):
        result = value.strftime("%Y-%m")
    else:
        result = str(value or "").strip()
    if not MONTH_PATTERN.fullmatch(result):
        raise WorkbookFormatError(f"{location} 的月份无效，应为 YYYY-MM。")
    month = int(result[5:7])
    if not 1 <= month <= 12:
        raise WorkbookFormatError(f"{location} 的月份无效：{result}。")
    return result


def _valid_global_key(value: Any, *, location: str) -> str:
    result = str(value or "").strip()
    if not result or len(result) > 240 or CONTROL_PATTERN.search(result):
        raise WorkbookFormatError(f"{location} 的记录类型无效。")
    return result


def _parse_payload(value: Any, *, location: str) -> dict[str, Any]:
    if not isinstance(value, str) or not value.strip():
        raise WorkbookFormatError(f"{location} 的 payload_json 为空。")
    if len(value) > MAX_JSON_CELL_CHARS:
        raise WorkbookFormatError(
            f"{location} 的 payload_json 超过 Excel 单元格 32,767 字符上限。"
        )
    if len(value.encode("utf-8")) > MAX_RECORD_JSON_BYTES:
        raise WorkbookFormatError(f"{location} 的 payload_json 过大。")

    def invalid_constant(token: str):
        raise ValueError(f"不允许的数值 {token}")

    try:
        payload = json.loads(value, parse_constant=invalid_constant)
        precise = json.loads(value, parse_float=Decimal, parse_constant=invalid_constant)
        payload = normalize_payload(exact_request_numbers(payload, precise))
    except (json.JSONDecodeError, ValueError, RecursionError) as exc:
        raise WorkbookFormatError(f"{location} 的 payload_json 已损坏：{exc}") from exc
    if not isinstance(payload, dict):
        raise WorkbookFormatError(f"{location} 的 payload_json 必须是 JSON 对象。")
    canonical_json(payload)
    return payload


def parse_internal_workbook(content: bytes) -> ParsedWorkbook:
    workbook = _open_workbook(content, data_only=False)
    try:
        missing_sheets = [name for name in INTERNAL_SHEETS if name not in workbook.sheetnames]
        if missing_sheets:
            raise WorkbookFormatError(
                f"这不是有效的内部 Excel，缺少工作表：{'、'.join(missing_sheets)}。"
            )

        _, meta_rows = _table_rows(
            workbook["__meta"], required=INTERNAL_SHEETS["__meta"]
        )
        meta_ids: set[str] = set()
        for row_number, row in meta_rows:
            if not any(value not in (None, "") for value in row.values()):
                continue
            meta_ids.add(
                _valid_dataset_id(
                    row.get("dataset_id"), location=f"__meta 第 {row_number} 行"
                )
            )
        if len(meta_ids) > 1:
            raise WorkbookFormatError("内部 Excel 的 __meta 同时包含多套 dataset_id。")

        records: list[WorkbookRecord] = []
        identities: set[tuple[str, str]] = set()
        record_dataset_ids: set[str] = set()
        warnings: list[str] = []
        mismatch_count = 0
        sheet_specs = (
            ("__day_records", "day", "日期"),
            ("__month_records", "month", "月份"),
            ("__global_records", "global", "记录类型"),
        )
        for sheet_name, record_type, key_column in sheet_specs:
            _, rows = _table_rows(
                workbook[sheet_name], required=INTERNAL_SHEETS[sheet_name]
            )
            for row_number, row in rows:
                location = f"{sheet_name} 第 {row_number} 行"
                source_dataset_id = _valid_dataset_id(
                    row.get("dataset_id"), location=location
                )
                record_dataset_ids.add(source_dataset_id)
                if record_type == "day":
                    key = _valid_day(row.get(key_column), location=location)
                elif record_type == "month":
                    key = _valid_month(row.get(key_column), location=location)
                else:
                    key = _valid_global_key(row.get(key_column), location=location)
                identity = (record_type, key)
                if identity in identities:
                    raise WorkbookFormatError(
                        f"内部 Excel 存在重复记录：{record_type}/{key}。"
                    )
                identities.add(identity)
                payload = normalize_fixed_expense_cash_flags(
                    _parse_payload(row.get("payload_json"), location=location)
                )
                checksum = payload_checksum(payload)
                supplied = str(row.get("sha256") or "").strip().lower()
                if supplied != checksum:
                    mismatch_count += 1
                    if len(warnings) < 20:
                        warnings.append(f"{location} 的 SHA-256 已按实际内容重算。")
                records.append(WorkbookRecord(record_type, key, payload, checksum))
                if len(records) > MAX_RECORDS:
                    raise WorkbookFormatError("内部 Excel 的记录数超过安全上限。")

        if len(record_dataset_ids) > 1:
            raise WorkbookFormatError("内部 Excel 同时包含多套 dataset_id，不能自动合并。")
        if records:
            if not meta_ids:
                raise WorkbookFormatError("内部 Excel 有业务记录，但 __meta 中没有 dataset_id。")
            if not record_dataset_ids.issubset(meta_ids):
                raise WorkbookFormatError("内部记录的 dataset_id 与 __meta 不一致。")
        source_ids = meta_ids | record_dataset_ids
        if len(source_ids) > 1:
            raise WorkbookFormatError("内部 Excel 的 dataset_id 不一致。")
        if mismatch_count > len(warnings):
            warnings.append(f"另有 {mismatch_count - len(warnings)} 条记录的 SHA-256 已重算。")
        return ParsedWorkbook(
            records=records,
            warnings=warnings,
            source_dataset_id=next(iter(source_ids), None),
            present_sheets=set(INTERNAL_SHEETS),
        )
    finally:
        workbook.close()


def _json_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _normalize_user_row(
    row: dict[str, Any], *, key_column: str | None, location: str, sheet_name: str
) -> dict[str, Any]:
    normalized = {key: _json_value(value) for key, value in row.items()}
    if key_column == "日期":
        normalized[key_column] = _valid_day(normalized.get(key_column), location=location)
    elif key_column == "月份":
        normalized[key_column] = _valid_month(normalized.get(key_column), location=location)
    try:
        normalized = normalize_payload({sheet_name: [normalized]})[sheet_name][0]
    except ValueError as exc:
        raise WorkbookFormatError(f"{location} 金额无效：{exc}") from exc
    canonical_json(normalized)
    return normalized


def parse_user_workbook(
    content: bytes,
    *,
    year: int,
    expected_dataset_id: str | None = None,
) -> ParsedWorkbook:
    if not 2000 <= int(year) <= 2200:
        raise WorkbookFormatError("year 必须在 2000 至 2200 之间。")
    workbook = _open_workbook(content, data_only=False)
    try:
        missing_sheets = [
            name for name in USER_SHEETS if name not in workbook.sheetnames
        ]
        if missing_sheets:
            raise WorkbookFormatError(
                "用户 Excel 将整体替换所选年份，为防止误删，必须包含完整六张工作表；"
                f"当前缺少：{'、'.join(missing_sheets)}。"
            )

        # A formula can be silently returned as an empty cached value when a
        # workbook is opened data_only.  User uploads are an authoritative
        # replacement, so reject formulas explicitly instead of importing an
        # incomplete snapshot whose result depends on a desktop Excel recalc.
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        raise WorkbookFormatError(
                            f"工作表「{sheet.title}」{cell.coordinate} 含有公式；"
                            "请先复制并粘贴为数值后再上传。"
                        )

        metadata_year: int | None = None
        source_dataset_id: str | None = None
        if "__app_meta" in workbook.sheetnames:
            _, metadata_rows = _table_rows(
                workbook["__app_meta"], required=("key", "value")
            )
            for row_number, row in metadata_rows:
                metadata_key = str(row.get("key") or "").strip()
                if metadata_key == "user_year":
                    try:
                        metadata_year = int(row.get("value"))
                    except (TypeError, ValueError) as exc:
                        raise WorkbookFormatError(
                            f"__app_meta 第 {row_number} 行的 user_year 无效。"
                        ) from exc
                elif metadata_key == "family_id":
                    raw_family_id = str(row.get("value") or "").strip()
                    if raw_family_id:
                        source_dataset_id = _valid_dataset_id(
                            raw_family_id,
                            location=f"__app_meta 第 {row_number} 行",
                        )
        if metadata_year is not None and metadata_year != int(year):
            raise WorkbookFormatError(
                f"用户 Excel 属于 {metadata_year} 年，不能按 {year} 年导入。"
            )
        if (
            expected_dataset_id is not None
            and source_dataset_id is not None
            and source_dataset_id != str(expected_dataset_id)
        ):
            raise WorkbookFormatError(
                f"用户 Excel 属于账本 {source_dataset_id}，不能替换当前账本 {expected_dataset_id}。"
            )
        present = {name for name in USER_SHEETS if name in workbook.sheetnames}
        tables: dict[str, list[dict[str, Any]]] = {}
        for sheet_name in USER_SHEETS:
            if sheet_name not in present:
                continue
            aliases = {"食材名称": "词条名称"} if sheet_name == "ingredients" else None
            key_column = DATE_SHEETS.get(sheet_name)
            if sheet_name == "monthly_fixed":
                key_column = "月份"
            # This upload is an authoritative replacement, not a best-effort
            # merge.  Requiring the complete canonical schema prevents a
            # damaged or hand-made workbook containing only key columns from
            # being accepted as an empty year and deleting valid cloud data.
            required = list(USER_SHEETS[sheet_name])
            _, raw_rows = _table_rows(
                workbook[sheet_name], required=required, aliases=aliases
            )
            rows: list[dict[str, Any]] = []
            for row_number, row in raw_rows:
                location = f"{sheet_name} 第 {row_number} 行"
                normalized = _normalize_user_row(
                    row,
                    key_column=key_column,
                    location=location,
                    sheet_name=sheet_name,
                )
                if sheet_name == "fixed_expenses":
                    # A manually-added old column must not turn a user Excel
                    # into a way to read or overwrite internal cash metadata.
                    normalized = {
                        key: value
                        for key, value in normalized.items()
                        if key not in INTERNAL_FIXED_EXPENSE_FIELDS
                    }
                if sheet_name == "ingredients" and not str(
                    normalized.get("词条名称") or ""
                ).strip():
                    raise WorkbookFormatError(f"{location} 的词条名称为空。")
                if sheet_name == "fixed_categories" and not str(
                    normalized.get("类别名称") or ""
                ).strip():
                    raise WorkbookFormatError(f"{location} 的类别名称为空。")
                rows.append(normalized)
            tables[sheet_name] = rows
        if sum(len(rows) for rows in tables.values()) > MAX_RECORDS:
            raise WorkbookFormatError("用户 Excel 的总记录数超过安全上限。")

        records: list[WorkbookRecord] = []
        warnings: list[str] = []
        target_prefix = f"{int(year):04d}-"
        workbook_years: set[int] = set()
        for sheet_name, key_column in DATE_SHEETS.items():
            workbook_years.update(
                int(str(row[key_column])[:4])
                for row in tables.get(sheet_name, [])
            )
        workbook_years.update(
            int(str(row["月份"])[:4]) for row in tables.get("monthly_fixed", [])
        )
        unexpected_years = workbook_years - {int(year)}
        if unexpected_years:
            found = "、".join(str(item) for item in sorted(unexpected_years))
            raise WorkbookFormatError(
                f"工作簿包含 {found} 年的数据；整年替换文件只能包含 {year} 年。"
            )
        present_date_sheets = [name for name in DATE_SHEETS if name in present]
        day_keys = sorted(
            {
                str(row[DATE_SHEETS[sheet_name]])
                for sheet_name in present_date_sheets
                for row in tables.get(sheet_name, [])
                if str(row[DATE_SHEETS[sheet_name]]).startswith(target_prefix)
            }
        )
        for day in day_keys:
            payload: dict[str, Any] = {}
            for sheet_name in present_date_sheets:
                key_column = DATE_SHEETS[sheet_name]
                payload[sheet_name] = [
                    row for row in tables.get(sheet_name, []) if row[key_column] == day
                ]
                payload[sheet_name].sort(key=canonical_json)
            records.append(
                WorkbookRecord("day", day, payload, payload_checksum(payload))
            )

        if "monthly_fixed" in present:
            months = sorted(
                {
                    str(row["月份"])
                    for row in tables.get("monthly_fixed", [])
                    if str(row["月份"]).startswith(target_prefix)
                }
            )
            for month in months:
                payload = {
                    "monthly_fixed": [
                        row
                        for row in tables.get("monthly_fixed", [])
                        if row["月份"] == month
                    ]
                }
                payload["monthly_fixed"].sort(key=canonical_json)
                records.append(
                    WorkbookRecord("month", month, payload, payload_checksum(payload))
                )

        for record_type in ("ingredients", "fixed_categories"):
            if record_type not in present:
                continue
            payload = {record_type: sorted(tables.get(record_type, []), key=canonical_json)}
            records.append(
                WorkbookRecord(
                    "global", record_type, payload, payload_checksum(payload)
                )
            )
        return ParsedWorkbook(
            records=records,
            warnings=warnings,
            source_dataset_id=source_dataset_id,
            present_sheets=present,
        )
    finally:
        workbook.close()


def _record_parts(record: Any) -> tuple[str, str, dict[str, Any], str]:
    if isinstance(record, dict):
        record_type = str(record.get("record_type") or "")
        record_key = str(record.get("record_key") or "")
        payload = record.get("payload") or {}
        checksum = str(record.get("checksum") or "")
    else:
        record_type = str(record.record_type)
        record_key = str(record.record_key)
        payload = record.payload or {}
        checksum = str(record.checksum or "")
    if not isinstance(payload, dict):
        raise WorkbookFormatError(f"云端记录 {record_type}/{record_key} 的 payload 不是对象。")
    return record_type, record_key, payload, checksum


def _safe_excel_value(value: Any, *, location: str) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    elif isinstance(value, (datetime, date, time)):
        value = value.isoformat()
    elif isinstance(value, float) and not math.isfinite(value):
        value = ""
    elif not isinstance(value, (str, int, float, bool)):
        value = str(value)
    if isinstance(value, str):
        if len(value) > MAX_CELL_CHARS:
            raise WorkbookFormatError(f"{location} 超过 Excel 单元格 32,767 字符上限。")
        if CONTROL_PATTERN.search(value):
            raise WorkbookFormatError(f"{location} 包含 Excel 不允许的控制字符。")
    return value


def _append_safe(sheet: Any, values: Iterable[Any], *, row_label: str) -> None:
    is_pristine = (
        sheet.max_row == 1
        and sheet.max_column == 1
        and sheet.cell(row=1, column=1).value is None
    )
    row_number = 1 if is_pristine else sheet.max_row + 1
    for column_number, raw in enumerate(values, start=1):
        value = _safe_excel_value(
            raw, location=f"{row_label} 第 {row_number} 行第 {column_number} 列"
        )
        cell = sheet.cell(row=row_number, column=column_number)
        cell.value = value
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"


def _style_table(sheet: Any) -> None:
    if sheet.max_row < 1:
        return
    fill = PatternFill("solid", fgColor="DCEAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="12304A")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    if sheet.max_column and sheet.max_row:
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for column_number in range(1, sheet.max_column + 1):
        width = 10
        for row_number in range(1, min(sheet.max_row, 200) + 1):
            value = sheet.cell(row=row_number, column=column_number).value
            width = max(width, min(36, len(str(value or "")) + 2))
        sheet.column_dimensions[get_column_letter(column_number)].width = width


def build_internal_workbook(
    *,
    dataset_id: str,
    schema_version: int,
    records: Iterable[Any],
    audit_records: Iterable[dict[str, Any]] = (),
) -> bytes:
    workbook = Workbook()
    info = workbook.active
    info.title = "内部数据说明"
    info["A1"] = "餐饮成本管家内部数据"
    info["A1"].font = Font(size=16, bold=True, color="12304A")
    info["A3"] = "此文件由系统生成，是同步与备份用的内部数据容器。请勿手工修改隐藏工作表。"
    info["A5"] = "导出时间"
    info["B5"] = datetime.now(timezone.utc).isoformat()
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 80

    sheets = {}
    for name, columns in INTERNAL_SHEETS.items():
        sheet = workbook.create_sheet(name)
        _append_safe(sheet, columns, row_label=name)
        sheets[name] = sheet
    audit_sheet = workbook.create_sheet(INTERNAL_AUDIT_SHEET)
    _append_safe(audit_sheet, INTERNAL_AUDIT_COLUMNS, row_label=INTERNAL_AUDIT_SHEET)

    ordered = sorted((_record_parts(item) for item in records), key=lambda item: (item[0], item[1]))
    for record_type, record_key, payload, _ in ordered:
        payload = normalize_fixed_expense_cash_flags(payload)
        payload_text = storage_json_dumps(payload, sort_keys=True)
        if len(payload_text) > MAX_JSON_CELL_CHARS:
            raise WorkbookFormatError(
                f"记录 {record_type}/{record_key} 的 JSON 超过 Excel 单元格 32,767 字符上限。"
            )
        checksum = hashlib.sha256(payload_text.encode("utf-8")).hexdigest()
        if record_type == "day":
            target, key = sheets["__day_records"], _valid_day(record_key, location="云端记录")
        elif record_type == "month":
            target, key = sheets["__month_records"], _valid_month(record_key, location="云端记录")
        elif record_type == "global":
            target, key = sheets["__global_records"], _valid_global_key(record_key, location="云端记录")
        else:
            continue
        _append_safe(target, (dataset_id, key, payload_text, checksum), row_label=target.title)

    # Modification history is cached only in the internal workbook.  It is not
    # a synchronised business record: imports intentionally ignore this
    # optional hidden sheet so a backup cannot replay or forge audit events.
    for audit in audit_records:
        if not isinstance(audit, dict):
            continue
        audit_id = str(audit.get("id") or "").strip()
        if not audit_id:
            continue
        payload_text = storage_json_dumps(audit, sort_keys=True)
        if len(payload_text) > MAX_JSON_CELL_CHARS:
            # The cache must never make a business-data backup unavailable.
            # Keep a compact audit card that is sufficient for the change-log
            # UI instead of attempting an unsafe partial JSON truncation.
            audit = {
                "id": audit_id,
                "dataset_id": audit.get("dataset_id", dataset_id),
                "created_at": audit.get("created_at") or audit.get("occurred_at") or "",
                "actor_username": audit.get("actor_username") or audit.get("actor_label") or "",
                "record_type": audit.get("record_type") or "",
                "record_key": audit.get("record_key") or "",
                "action": audit.get("action") or "",
                "change_summary": audit.get("change_summary") or "修改记录内容过长",
            }
            payload_text = storage_json_dumps(audit, sort_keys=True)
        if len(payload_text) > MAX_JSON_CELL_CHARS:
            # Even a malformed/huge summary should not block export.  The
            # server remains the authoritative source for this optional cache.
            continue
        checksum = hashlib.sha256(payload_text.encode("utf-8")).hexdigest()
        _append_safe(
            audit_sheet,
            (
                dataset_id,
                audit_id,
                str(audit.get("created_at") or audit.get("occurred_at") or ""),
                payload_text,
                checksum,
            ),
            row_label=INTERNAL_AUDIT_SHEET,
        )

    _append_safe(
        sheets["__meta"],
        (dataset_id, int(schema_version), datetime.now(timezone.utc).isoformat()),
        row_label="__meta",
    )
    for sheet in sheets.values():
        _style_table(sheet)
        sheet.sheet_state = "veryHidden"
    _style_table(audit_sheet)
    audit_sheet.sheet_state = "veryHidden"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _payload_rows(
    payload: dict[str, Any],
    field: str,
    *,
    key: str | None = None,
    key_column: str | None = None,
) -> list[dict[str, Any]]:
    raw = payload.get(field, [])
    if raw in (None, ""):
        return []
    if not isinstance(raw, list):
        raise WorkbookFormatError(f"记录 {key} 的 {field} 不是数组。")
    rows: list[dict[str, Any]] = []
    for item in raw:
        if not isinstance(item, dict):
            raise WorkbookFormatError(f"记录 {key} 的 {field} 含有非对象行。")
        row = dict(item)
        if key_column is not None and key is not None:
            row[key_column] = key
        rows.append(row)
    return rows


def _ordered_headers(sheet_name: str, rows: list[dict[str, Any]]) -> list[str]:
    headers = list(USER_SHEETS[sheet_name])
    seen = set(headers)
    for row in rows:
        for key in row:
            key = str(key)
            if key and key not in seen:
                seen.add(key)
                headers.append(key)
    return headers


def build_user_workbook(
    *,
    dataset_id: str,
    schema_version: int,
    year: int,
    records: Iterable[Any],
) -> bytes:
    if not 2000 <= int(year) <= 2200:
        raise WorkbookFormatError("year 必须在 2000 至 2200 之间。")
    collected: dict[str, list[dict[str, Any]]] = {name: [] for name in USER_SHEETS}
    prefix = f"{int(year):04d}-"
    for record_type, record_key, payload, _ in sorted(
        (_record_parts(item) for item in records), key=lambda item: (item[0], item[1])
    ):
        if record_type == "day" and record_key.startswith(prefix):
            valid_key = _valid_day(record_key, location="云端记录")
            for sheet_name in DATE_SHEETS:
                rows = _payload_rows(
                    payload,
                    sheet_name,
                    key=valid_key,
                    key_column=DATE_SHEETS[sheet_name],
                )
                if sheet_name == "fixed_expenses":
                    rows = [
                        {
                            field: value
                            for field, value in row.items()
                            if field not in INTERNAL_FIXED_EXPENSE_FIELDS
                        }
                        for row in rows
                    ]
                collected[sheet_name].extend(rows)
        elif record_type == "month" and record_key.startswith(prefix):
            valid_key = _valid_month(record_key, location="云端记录")
            collected["monthly_fixed"].extend(
                _payload_rows(
                    payload,
                    "monthly_fixed",
                    key=valid_key,
                    key_column="月份",
                )
            )
        elif record_type == "global" and record_key in {"ingredients", "fixed_categories"}:
            collected[record_key].extend(
                _payload_rows(
                    payload,
                    record_key,
                    key=record_key,
                )
            )

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in USER_SHEETS:
        # Export monetary cells as decimal strings. A workbook may be opened
        # and resaved by Excel without converting exact cents to binary floats.
        try:
            rows = normalize_payload({sheet_name: collected[sheet_name]})[sheet_name]
        except ValueError as exc:
            raise WorkbookFormatError(f"{sheet_name} 金额无效：{exc}") from exc
        if sheet_name in DATE_SHEETS:
            rows.sort(
                key=lambda row: (
                    str(row.get(DATE_SHEETS[sheet_name], "")),
                    canonical_json(row),
                )
            )
        elif sheet_name == "monthly_fixed":
            rows.sort(
                key=lambda row: (str(row.get("月份", "")), canonical_json(row))
            )
        else:
            rows.sort(key=canonical_json)
        headers = _ordered_headers(sheet_name, rows)
        sheet = workbook.create_sheet(sheet_name)
        _append_safe(sheet, headers, row_label=sheet_name)
        for row in rows:
            _append_safe(sheet, (row.get(header, "") for header in headers), row_label=sheet_name)
        _style_table(sheet)

    meta = workbook.create_sheet("__app_meta")
    _append_safe(meta, ("key", "value"), row_label="__app_meta")
    _append_safe(meta, ("family_id", dataset_id), row_label="__app_meta")
    _append_safe(meta, ("user_year", int(year)), row_label="__app_meta")
    _append_safe(meta, ("schema_version", int(schema_version)), row_label="__app_meta")
    meta.sheet_state = "veryHidden"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
